"""The WorksheetToolkit class.

Each method formats part of one worksheet and returns the toolkit, so calls can be
chained. Picking cells, reading colors and measuring text are handled by the
private modules alongside this one.
"""

from __future__ import annotations

from collections.abc import Iterable
from copy import copy
from typing import cast, get_args

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.cell_range import CellRange
from openpyxl.worksheet.views import Selection
from openpyxl.worksheet.worksheet import Worksheet

from ._colors import has_color, normalize_color
from ._limits import MAX_COLUMN_WIDTH, MAX_ROW_HEIGHT
from ._ranges import (
    as_indexes,
    check_bounds,
    iter_cells,
    resolve_cells,
    resolve_range_arguments,
)
from ._sentinel import UNCHANGED, Unchanged
from ._text import displayed_text, measure_text, workbook_normal_font
from ._types import (
    BorderSide,
    BorderStyle,
    FillType,
    HorizontalAlignment,
    IndexSelection,
    ReadingOrder,
    Underline,
    VerticalAlignment,
)

__all__ = ["WorksheetToolkit"]

# One list, so the runtime check and the type a caller sees cannot drift apart.
_BORDER_SIDES: tuple[BorderSide, ...] = get_args(BorderSide)


class WorksheetToolkit:
    """Formatting helpers for one openpyxl worksheet.

    Each method merges what it is given into the formatting a cell already
    carries, rather than replacing it, and returns the toolkit so that calls can
    be chained. A parameter left out keeps the value it had; None is a value in
    its own right, and clears the thing it names.

    Examples
    --------
    >>> toolkit.set_font(cells="A1:C1", bold=True).set_fill(
    ...     cells="A1:C1", fill_type='solid', start_color='#eef2f6'
    ... )
    """

    def __init__(self, worksheet: Worksheet) -> None:
        """Initialize a WorksheetToolkit instance for a single worksheet.

        Parameters
        ----------
        worksheet : openpyxl.worksheet.worksheet.Worksheet
            The openpyxl Worksheet object to operate on. All formatting and
            utility methods in this instance will apply to this worksheet.

        Raises
        ------
        TypeError
            If given anything other than a worksheet.
        """
        if not isinstance(worksheet, Worksheet):
            raise TypeError(
                f"expected a Worksheet, got {type(worksheet).__name__}. "
                f"A workbook's sheet is wb['Sheet name'] or wb.active."
            )
        self.worksheet = worksheet

    def __repr__(self) -> str:
        """Name of the sheet being worked on."""
        return f"<{type(self).__name__} {self.worksheet.title!r}>"

    def freeze_panes(self, cell: str | None = None) -> WorksheetToolkit:
        """Freeze the rows above and the columns left of a cell. Those rows and columns
        then stay in view as the sheet is scrolled.

        Parameters
        ----------
        cell : str, optional
            The top-left cell of the scrolling area, such as 'B2'. None or an empty
            string unfreezes the panes.

        Returns
        -------
        WorksheetToolkit
        """
        if cell is None or cell == "":
            # openpyxl clears the pane but leaves the three split selections behind,
            # which is what makes some versions of Excel offer to repair the file.
            self.worksheet.freeze_panes = None
            self.worksheet.sheet_view.selection = [Selection()]
        else:
            self.worksheet.freeze_panes = cell
        return self

    def merge_cells(
        self,
        *,
        cells: str | None = None,
        start_row: int | None = None,
        start_column: int | None = None,
        end_row: int | None = None,
        end_column: int | None = None,
    ) -> WorksheetToolkit:
        """Merge a rectangular range of cells into one.

        The merged cell takes the value of the top-left cell; the rest are cleared.
        Merging a range that is already merged does nothing.

        Parameters
        ----------
        cells : str, optional
            The range, such as 'A1:C3'. Give this or all four coordinates, not both.
        start_row, start_column, end_row, end_column : int, optional
            The top-left and bottom-right corners of the block.

        Returns
        -------
        WorksheetToolkit

        Raises
        ------
        ValueError
            If both ``cells`` and the coordinates are given, or if only some of the four
            coordinates are.
        """
        target = resolve_range_arguments(
            "merge_cells", cells, start_row, start_column, end_row, end_column
        )
        if isinstance(target, str):
            self.worksheet.merge_cells(range_string=target)
        else:
            first_row, first_column, last_row, last_column = target
            self.worksheet.merge_cells(
                start_row=first_row,
                start_column=first_column,
                end_row=last_row,
                end_column=last_column,
            )
        return self

    def unmerge_cells(
        self,
        *,
        cells: str | None = None,
        start_row: int | None = None,
        start_column: int | None = None,
        end_row: int | None = None,
        end_column: int | None = None,
    ) -> WorksheetToolkit:
        """Undo a merge, taking the same arguments as ``merge_cells``.

        A range that is not merged is left alone. A range that cannot be read at all
        raises.

        Parameters
        ----------
        cells : str, optional
            The range, such as 'A1:C3'. Give this or all four coordinates, not both.
        start_row, start_column, end_row, end_column : int, optional
            The top-left and bottom-right corners of the block.

        Returns
        -------
        WorksheetToolkit

        Raises
        ------
        ValueError
            If both ``cells`` and the coordinates are given, or if only some of the four
            coordinates are.
        """
        target = resolve_range_arguments(
            "unmerge_cells", cells, start_row, start_column, end_row, end_column
        )
        if isinstance(target, str):
            if CellRange(target) not in self.worksheet.merged_cells.ranges:
                return self
            self.worksheet.unmerge_cells(range_string=target)
        else:
            first_row, first_column, last_row, last_column = target
            block = CellRange(
                min_col=first_column,
                min_row=first_row,
                max_col=last_column,
                max_row=last_row,
            )
            if block not in self.worksheet.merged_cells.ranges:
                return self
            self.worksheet.unmerge_cells(
                start_row=first_row,
                start_column=first_column,
                end_row=last_row,
                end_column=last_column,
            )
        return self

    def set_alignment(
        self,
        *,
        cells: str | None = None,
        rows: IndexSelection = None,
        columns: IndexSelection = None,
        intersections_only: bool = True,
        horizontal: HorizontalAlignment | None | Unchanged = UNCHANGED,
        vertical: VerticalAlignment | None | Unchanged = UNCHANGED,
        text_rotation: int | None | Unchanged = UNCHANGED,
        wrap_text: bool | None | Unchanged = UNCHANGED,
        shrink_to_fit: bool | None | Unchanged = UNCHANGED,
        indent: float | None | Unchanged = UNCHANGED,
        reading_order: ReadingOrder | None | Unchanged = UNCHANGED,
    ) -> WorksheetToolkit:
        """Set how text sits inside its cells.

        Parameters
        ----------
        cells : str, optional
            The cells to format, as a range: a block such as 'A1:C3', whole columns
            such as 'B:D', whole rows such as '2:5', or one cell such as 'C3'. Give
            this or ``rows`` and ``columns``, not both.
        rows : list of int, optional
            Row numbers. Defaults to every row in use.
        columns : list of int, optional
            Column numbers. Defaults to every column in use.
        intersections_only : bool, optional
            True, the default, formats the cells where the given rows and the given
            columns cross. False formats every cell in those rows and every cell in
            those columns, which is a cross rather than a block. Only considered when
            both ``rows`` and ``columns`` are given.
        horizontal : str, optional
            'general', 'left', 'center', 'right', 'fill', 'justify',
            'centerContinuous' or 'distributed'.
        vertical : str, optional
            'top', 'center', 'bottom', 'justify' or 'distributed'.
        text_rotation : int, optional
            Degrees, 0 to 180. 255 is the separate case: letters stacked one above the
            next, reading downwards. Nothing between 181 and 254 is accepted.
        wrap_text : bool, optional
            True breaks long text onto more than one line inside the cell.
        shrink_to_fit : bool, optional
            True reduces the displayed font size until the text fits the column.
        indent : int, optional
            Indent levels, not spaces. Excel renders one level as roughly three
            characters of the normal font. 0 to 255, and None is another way of
            saying 0, which is no indent.
        reading_order : int, optional
            Text direction: 0 leaves it to the text, 1 is left-to-right, 2 is
            right-to-left. None is another way of saying 0.

        Returns
        -------
        WorksheetToolkit

        Raises
        ------
        TypeError
            If ``rows`` or ``columns`` is given a single number rather than a
            list, or an index that is not an integer.
        ValueError
            If both ``cells`` and ``rows`` or ``columns`` are given, if ``cells``
            cannot be read as a range, or if an index falls outside Excel's grid.

        Examples
        --------
        Center every cell, both ways:

        >>> toolkit.set_alignment(horizontal='center', vertical='center')

        Wrap the text in rows 1 and 2 only:

        >>> toolkit.set_alignment(rows=[1, 2], wrap_text=True)

        Rotate the four cells where rows 1 and 2 meet columns 1 and 2:

        >>> toolkit.set_alignment(
        ...     rows=[1, 2], columns=[1, 2], intersections_only=True, text_rotation=45
        ... )
        """
        parameters = {
            "horizontal": horizontal,
            "vertical": vertical,
            "text_rotation": text_rotation,
            "wrap_text": wrap_text,
            "shrink_to_fit": shrink_to_fit,
            "indent": 0 if indent is None else indent,
            "readingOrder": 0 if reading_order is None else reading_order,
        }

        updates = []
        for cell in iter_cells(self.worksheet, rows, columns, intersections_only, cells):
            # Copy and override for the same reason as set_font: rebuilding from
            # the arguments drops justifyLastLine and relativeIndent.
            new_alignment = cast(Alignment, copy(cell.alignment))
            for attribute, value in parameters.items():
                if value is not UNCHANGED:
                    setattr(new_alignment, attribute, value)
            updates.append((cell, new_alignment))

        for cell, alignment in updates:
            cell.alignment = alignment
        return self

    def set_border(
        self,
        *,
        cells: str | None = None,
        rows: IndexSelection = None,
        columns: IndexSelection = None,
        intersections_only: bool = True,
        sides: BorderSide | Iterable[BorderSide] = (
            "top",
            "bottom",
            "left",
            "right",
        ),
        style: BorderStyle | None | Unchanged = UNCHANGED,
        color: str | None | Unchanged = UNCHANGED,
    ) -> WorksheetToolkit:
        """Set the border on one or more sides of each cell.

        Sides that are not named keep the border they had. A color on its own cannot
        be applied to a side with no line.

        Parameters
        ----------
        cells : str, optional
            The cells to format, as a range: a block such as 'A1:C3', whole columns
            such as 'B:D', whole rows such as '2:5', or one cell such as 'C3'. Give
            this or ``rows`` and ``columns``, not both.
        rows : list of int, optional
            Row numbers. Defaults to every row in use.
        columns : list of int, optional
            Column numbers. Defaults to every column in use.
        intersections_only : bool, optional
            True, the default, formats the cells where the given rows and the given
            columns cross. False formats every cell in those rows and every cell in
            those columns, which is a cross rather than a block. Only considered when
            both ``rows`` and ``columns`` are given.
        sides : str or list of str, optional
            Which sides to change: 'top', 'bottom', 'left', 'right', 'diagonal_up' or
            'diagonal_down'. Defaults to the four straight sides. One side may be
            given on its own, without a list.
        style : str, optional
            'hair', 'thin', 'medium', 'thick', 'double', 'dotted', 'dashed',
            'mediumDashed', 'mediumDashDot', 'mediumDashDotDot', 'dashDot',
            'dashDotDot' or 'slantDashDot'. None removes the line, and on a diagonal
            takes that direction away.
        color : str, optional
            Hex color code, such as '#ff0000'.

        Returns
        -------
        WorksheetToolkit

        Raises
        ------
        TypeError
            If ``rows`` or ``columns`` is given a single number rather than a
            list, or an index that is not an integer.
        ValueError
            If a side is not one of the six names, if a color is given for a side
            with no line and no ``style`` to draw one, if both ``cells`` and
            ``rows`` or ``columns`` are given, or if an index falls outside
            Excel's grid.

        Notes
        -----
        A cell stores one diagonal line and a flag for each direction, so the two
        diagonals cannot carry different styles or colors: styling one while the
        other is drawn restyles both. Taking one away leaves the other as it was.

        Examples
        --------
        Thin red rules above and below rows 1 and 2:

        >>> toolkit.set_border(
        ...     rows=[1, 2], sides=('top', 'bottom'), style='thin', color='#ff0000'
        ... )

        A diagonal through B2, bottom-left to top-right. ``sides`` takes one name
        or several, so a single side needs no comma:

        >>> toolkit.set_border(
        ...     rows=[2], columns=[2], sides='diagonal_up',
        ...     style='thin', color='#00ff00',
        ... )
        """
        if isinstance(sides, str):
            sides = (sides,)
        unknown = [side for side in sides if side not in _BORDER_SIDES]
        if unknown:
            raise ValueError(f"unknown border side(s) {unknown}. Choose from {list(_BORDER_SIDES)}")

        # Built first and assigned afterwards, so a cell that cannot be given the
        # requested border does not leave the rest of the range half-drawn.
        updates = []
        for cell in iter_cells(self.worksheet, rows, columns, intersections_only, cells):
            current = cell.border
            straight: dict[str, Side] = {}

            # Standard sides
            for side_name in ("left", "right", "top", "bottom"):
                if side_name in sides:
                    new_style = (
                        style if style is not UNCHANGED else getattr(current, side_name).style
                    )
                    if color is not UNCHANGED and new_style is None:
                        raise ValueError(
                            f"{cell.coordinate} has no {side_name} border, so a color on its "
                            f"own would not show. Pass style='thin' as well."
                        )
                    straight[side_name] = Side(
                        style=new_style,
                        color=normalize_color(color)
                        if color is not UNCHANGED
                        else getattr(current, side_name).color,
                    )
                else:
                    straight[side_name] = getattr(current, side_name)

            # Diagonal side logic
            if "diagonal_up" in sides or "diagonal_down" in sides:
                new_style = style if style is not UNCHANGED else current.diagonal.style
                if color is not UNCHANGED and new_style is None:
                    raise ValueError(
                        f"{cell.coordinate} has no diagonal border, so a color on its own "
                        f"would not show. Pass style='thin' as well."
                    )
                # A cell holds one diagonal line and a flag per direction, so a
                # direction is drawn only where its flag and that line agree. Only
                # the direction actually named is touched: rewriting both flags
                # every time is what made one call clear the other.
                drawn = new_style is not None
                diagonal_up = drawn if "diagonal_up" in sides else current.diagonalUp
                diagonal_down = drawn if "diagonal_down" in sides else current.diagonalDown
                if not diagonal_up and not diagonal_down:
                    # Neither direction is left, so the line they shared goes too.
                    diagonal = Side()
                else:
                    diagonal = Side(
                        # Taking one direction away leaves the other drawing this
                        # line, so the line itself is only rewritten when asked for.
                        style=new_style if drawn else current.diagonal.style,
                        color=normalize_color(color)
                        if color is not UNCHANGED
                        else current.diagonal.color,
                    )
            else:
                diagonal = current.diagonal
                diagonal_up = current.diagonalUp
                diagonal_down = current.diagonalDown

            updates.append(
                (
                    cell,
                    Border(
                        left=straight["left"],
                        right=straight["right"],
                        top=straight["top"],
                        bottom=straight["bottom"],
                        diagonal=diagonal,
                        diagonalUp=diagonal_up,
                        diagonalDown=diagonal_down,
                    ),
                )
            )

        for cell, border in updates:
            cell.border = border

        return self

    def set_outside_border(
        self,
        *,
        style: BorderStyle | None,
        cells: str | None = None,
        start_row: int | None = None,
        end_row: int | None = None,
        start_column: int | None = None,
        end_column: int | None = None,
        color: str | None | Unchanged = UNCHANGED,
    ) -> WorksheetToolkit:
        """Draw a border around the edge of a block, leaving the inside alone.

        Parameters
        ----------
        style : str
            Border style, as in ``set_border``. Required.
        cells : str, optional
            The range, such as 'A1:C3'. Give this or all four coordinates, not both.
        start_row, start_column, end_row, end_column : int, optional
            The top-left and bottom-right corners of the block.
        color : str, optional
            Hex color code, such as '#000000'.

        Returns
        -------
        WorksheetToolkit

        Raises
        ------
        ValueError
            If both ``cells`` and the coordinates are given, if only some of the four
            coordinates are, or if ``cells`` cannot be read as a range.

        Examples
        --------
        A thin black box around rows 1 to 3 and columns 1 to 4:

        >>> toolkit.set_outside_border(
        ...     start_row=1, end_row=3, start_column=1, end_column=4,
        ...     style='thin', color='#000000',
        ... )

        The same block, named as a range:

        >>> toolkit.set_outside_border(cells="A1:D3", style='medium')
        """
        target = resolve_range_arguments(
            "set_outside_border", cells, start_row, start_column, end_row, end_column
        )
        if isinstance(target, str):
            rows, columns = resolve_cells(self.worksheet, target)
        else:
            first_row, first_column, last_row, last_column = target
            rows = list(range(first_row, last_row + 1))
            columns = list(range(first_column, last_column + 1))
        start_row, end_row = rows[0], rows[-1]
        start_column, end_column = columns[0], columns[-1]

        # Top border
        self.set_border(
            rows=[start_row],
            columns=columns,
            intersections_only=True,
            sides=("top",),
            style=style,
            color=color,
        )
        # Bottom border
        self.set_border(
            rows=[end_row],
            columns=columns,
            intersections_only=True,
            sides=("bottom",),
            style=style,
            color=color,
        )
        # Left border
        self.set_border(
            rows=rows,
            columns=[start_column],
            intersections_only=True,
            sides=("left",),
            style=style,
            color=color,
        )
        # Right border
        self.set_border(
            rows=rows,
            columns=[end_column],
            intersections_only=True,
            sides=("right",),
            style=style,
            color=color,
        )

        return self

    def set_column_best_fit(
        self,
        *,
        columns: IndexSelection = None,
        padding: float = 0,
        ignore_rows: Iterable[int] | None = None,
        ignore_formulas: bool = True,
        ignore_wrapped: bool = True,
        ignore_merged: bool = True,
        min_width: float | None = None,
        max_width: float | None = None,
    ) -> WorksheetToolkit:
        """Widen each column to fit its widest cell.

        Every character is measured at its own width in the font being used. Excel's
        own formula turns the total into a column width:

            width = (pixels of text + 5 padding pixels) / max digit width

        A column with nothing in it is left alone rather than shrunk.

        Parameters
        ----------
        columns : list of int, optional
            Column numbers to fit. Defaults to every column in use.
        padding : float, optional
            Extra width on top of the fitted value. Defaults to 0; Excel's own 5
            pixels of cell padding are already part of the formula.
        ignore_rows : list of int, optional
            Row numbers to leave out when measuring, for a judgement this method
            cannot make on its own: a long unmerged title, a totals row with a
            wordy label, one outlier that would double the column.
        ignore_formulas : bool, optional
            True, the default, leaves formula cells unmeasured. The formula text is
            not what the reader sees, so measuring it oversizes the column.
        ignore_wrapped : bool, optional
            True, the default, leaves cells with ``wrap_text`` unmeasured. False
            widens the column to fit the text, resulting in cell contents on a single
            line.
        ignore_merged : bool, optional
            True, the default, leaves a cell merged across columns unmeasured.
            False measures it, widening the one column that holds the value to
            fit text the reader sees spread across the merge. A merge running
            down a single column is measured either way.
        min_width : float, optional
            Lower bound on the result.
        max_width : float, optional
            Upper bound. Defaults to Excel's own maximum of 255.

        Returns
        -------
        WorksheetToolkit

        Raises
        ------
        TypeError
            If ``columns`` is given a single number rather than a list, or an index that
            is not an integer.
        ValueError
            If a column falls outside Excel's grid.

        Notes
        -----
        Only dates and times are rendered as Excel displays them. Other number formats
        are measured as the value is stored, so a currency or percentage column may
        come out narrower than it needs to be.

        Built-in metrics cover Aptos, Arial, Calibri, Cambria, Courier New, Futura,
        Garamond, Georgia, Inter, Open Sans, Palatino, Roboto, Segoe UI, Tahoma, Times
        New Roman, and Verdana, along with Helvetica, Book Antiqua and the open faces
        drawn to match them. Any other face is measured with Verdana's character
        widths. Verdana is the widest of them so an unmeasurable face errs wide; a
        column that comes out wrong can be set directly with ``set_column_width``.
        """
        ws = self.worksheet
        if columns is None:
            columns = list(range(1, ws.max_column + 1))
        else:
            columns = as_indexes(columns, "columns")
        # Before any ws.cell() call: materialising an out-of-grid cell makes the
        # workbook permanently unsaveable, so a later failure would come too late.
        check_bounds(columns=columns)
        ignore_rows = set() if ignore_rows is None else set(ignore_rows)
        normal_font = workbook_normal_font(self.worksheet)
        # Read once rather than per cell. Only the top-left of a merge holds the
        # value; the rest are MergedCells with none, and are skipped anyway.
        spanning_merges = (
            {(r.min_row, r.min_col) for r in ws.merged_cells.ranges if r.max_col > r.min_col}
            if ignore_merged
            else set()
        )

        for col in columns:
            excel_width = 0.0
            measured_anything = False
            for row in range(1, ws.max_row + 1):
                if row in ignore_rows:
                    continue

                cell = ws.cell(row=row, column=col)

                if ignore_formulas and cell.data_type == "f":
                    continue

                if cell.value is None:
                    continue

                if (row, col) in spanning_merges:
                    continue

                if ignore_wrapped and cell.alignment.wrap_text:
                    continue

                measured_anything = True
                # A cell that inherits the workbook font reports neither name nor
                # size of its own, which Font(bold=True) alone is enough to produce.
                font = (
                    cell.font.name or normal_font[0],
                    normal_font[1] if cell.font.sz is None else cell.font.sz,
                )
                excel_width = max(
                    measure_text(displayed_text(cell), font, normal_font), excel_width
                )

            if not measured_anything:
                # Nothing to fit. Leaving the column alone matters because a width
                # set deliberately beforehand would otherwise be cut to the padding.
                continue

            width = excel_width + padding
            if min_width is not None:
                width = max(width, min_width)
            ws.column_dimensions[get_column_letter(col)].width = min(
                width, MAX_COLUMN_WIDTH if max_width is None else max_width
            )

        return self

    def set_column_width(self, *, width: float, columns: IndexSelection = None) -> WorksheetToolkit:
        """Set the width of one or more columns.

        Parameters
        ----------
        width : float
            Column width in Excel character units, not pixels. Required: there is no
            existing value to leave alone, so omitting it cannot mean anything. A
            width of 0 hides the column.
        columns : list of int, optional
            Column numbers to modify. Defaults to every column in use.

        Returns
        -------
        WorksheetToolkit

        Raises
        ------
        TypeError
            If ``columns`` is given a single number rather than a list, or an index that
            is not an integer.
        ValueError
            If ``width`` is negative, or if a column falls outside Excel's grid.
        """
        if width < 0:
            raise ValueError(f"width must not be negative, got {width}")

        if columns is None:
            columns = list(range(1, self.worksheet.max_column + 1))
        else:
            columns = as_indexes(columns, "columns")
        check_bounds(columns=columns)
        width = min(width, MAX_COLUMN_WIDTH)
        for col in columns:
            # get_column_letter rather than a cell lookup: row 1 of the column
            # may be a MergedCell, which has no column_letter at all.
            dimension = self.worksheet.column_dimensions[get_column_letter(col)]
            # openpyxl cannot persist a zero width -- the writer drops any falsy
            # dimension -- so the only way to honour it is to hide the column.
            if width == 0:
                dimension.hidden = True
            else:
                dimension.width = width
        return self

    def set_row_height(self, *, height: float, rows: IndexSelection = None) -> WorksheetToolkit:
        """Set the height of one or more rows.

        Parameters
        ----------
        height : float
            Row height in points. Required, for the same reason as the column width.
            A height of 0 hides the row.
        rows : list of int, optional
            Row numbers to modify. Defaults to every row in use.

        Returns
        -------
        WorksheetToolkit

        Raises
        ------
        TypeError
            If ``rows`` is given a single number rather than a list, or an index that is
            not an integer.
        ValueError
            If ``height`` is negative, or if a row falls outside Excel's grid.
        """
        if height < 0:
            raise ValueError(f"height must not be negative, got {height}")

        if rows is None:
            rows = list(range(1, self.worksheet.max_row + 1))
        else:
            rows = as_indexes(rows, "rows")
        check_bounds(rows=rows)
        height = min(height, MAX_ROW_HEIGHT)
        for row in rows:
            dimension = self.worksheet.row_dimensions[row]
            if height == 0:
                dimension.hidden = True
            else:
                dimension.height = height
        return self

    def set_fill(
        self,
        *,
        cells: str | None = None,
        rows: IndexSelection = None,
        columns: IndexSelection = None,
        intersections_only: bool = True,
        fill_type: FillType | None | Unchanged = UNCHANGED,
        start_color: str | Unchanged = UNCHANGED,
        end_color: str | Unchanged = UNCHANGED,
    ) -> WorksheetToolkit:
        """Set the background fill of cells.

        A fill has two parts: a pattern and the colors it is drawn in. A color with
        no pattern has nothing to show through, so a cell with no fill yet needs both.

        Parameters
        ----------
        cells : str, optional
            The cells to format, as a range: a block such as 'A1:C3', whole columns
            such as 'B:D', whole rows such as '2:5', or one cell such as 'C3'. Give
            this or ``rows`` and ``columns``, not both.
        rows : list of int, optional
            Row numbers. Defaults to every row in use.
        columns : list of int, optional
            Column numbers. Defaults to every column in use.
        intersections_only : bool, optional
            True, the default, formats the cells where the given rows and the given
            columns cross. False formats every cell in those rows and every cell in
            those columns, which is a cross rather than a block. Only considered when
            both ``rows`` and ``columns`` are given.
        fill_type : str, optional
            The pattern. 'solid' is the common one; the rest are hatches and shades
            such as 'gray125', 'lightGrid' and 'darkTrellis'. None removes the fill.
        start_color : str, optional
            Hex color code for the foreground. For a solid fill this is the color
            that shows.
        end_color : str, optional
            Hex color code for the background, which only a patterned fill draws.
            Excel ignores it for a solid fill.

        Returns
        -------
        WorksheetToolkit

        Raises
        ------
        TypeError
            If ``rows`` or ``columns`` is given a single number rather than a
            list, or an index that is not an integer.
        ValueError
            If a color is given as None, which cannot clear a fill that always
            carries one; if a color is given for a cell with no pattern and no
            ``fill_type`` to make one; if a ``fill_type`` is given for a cell
            with no color, which would paint it black; if both ``cells`` and
            ``rows`` or ``columns`` are given; or if an index falls outside
            Excel's grid.

        Examples
        --------
        A color needs a pattern to show through, so a cell with no fill yet takes
        both:

        >>> toolkit.set_fill(rows=[1], fill_type='solid', start_color='#f4d2d3')

        Once a cell has a pattern, the color can be changed on its own:

        >>> toolkit.set_fill(rows=[1], start_color='#ffff00')

        Yellow where rows 1 and 2 meet columns 1 and 2:

        >>> toolkit.set_fill(
        ...     rows=[1, 2], columns=[1, 2], intersections_only=True,
        ...     fill_type='solid', start_color='#ffff00',
        ... )
        """
        if start_color is None or end_color is None:
            # A pattern fill has no colorless state: its default foreground is an
            # opaque black, so None here would repaint rather than clear.
            raise ValueError(
                "set_fill cannot clear a color on its own -- a pattern fill always carries "
                "one. Pass fill_type=None to remove the fill, or give an explicit hex color."
            )

        requested = any(arg is not UNCHANGED for arg in (fill_type, start_color, end_color))

        # Every fill is built before any is assigned, so a failure part-way through
        # leaves the worksheet exactly as it was found rather than half-formatted.
        updates = []
        for cell in iter_cells(self.worksheet, rows, columns, intersections_only, cells):
            current = cell.fill
            # cell.fill is a StyleProxy, so isinstance against PatternFill never
            # matches; the wrapped object's tagname is what distinguishes them.
            if getattr(current, "tagname", None) == "patternFill":
                current_type = current.fill_type
                # copy(), not the object itself: a Color is stored by reference, so
                # passing the existing one through would leave this cell sharing a
                # mutable Color with whatever it was read from -- including the
                # process-global default that PatternFill uses for an absent bgColor.
                current_start = copy(current.start_color)
                current_end = copy(current.end_color)
            elif not requested:
                # Nothing was asked for, so leave the gradient alone rather than
                # flattening it into a blank pattern fill.
                continue
            else:
                # A GradientFill has no pattern or start/end color to merge with,
                # so anything left unspecified falls back to PatternFill's default.
                current_type = current_start = current_end = None

            new_type = current_type if fill_type is UNCHANGED else fill_type
            new_start = current_start if start_color is UNCHANGED else normalize_color(start_color)

            if start_color is not UNCHANGED and new_type is None:
                raise ValueError(
                    f"{cell.coordinate} has no fill pattern, so a color on its own would "
                    f"not show. Pass fill_type='solid' as well."
                )
            if fill_type not in (UNCHANGED, None) and not has_color(new_start):
                raise ValueError(
                    f"{cell.coordinate} has no fill color, so fill_type={fill_type!r} alone "
                    f"would paint it black. Pass start_color as well."
                )

            updates.append(
                (
                    cell,
                    PatternFill(
                        fill_type=new_type,
                        # The Color object is passed through rather than its .rgb:
                        # for a theme, indexed or automatic color that attribute is
                        # the descriptor itself, which PatternFill rejects.
                        start_color=new_start,
                        end_color=current_end
                        if end_color is UNCHANGED
                        else normalize_color(end_color),
                    ),
                )
            )

        for cell, fill in updates:
            cell.fill = fill
        return self

    def set_font(
        self,
        *,
        cells: str | None = None,
        rows: IndexSelection = None,
        columns: IndexSelection = None,
        intersections_only: bool = True,
        name: str | None | Unchanged = UNCHANGED,
        size: float | None | Unchanged = UNCHANGED,
        bold: bool | None | Unchanged = UNCHANGED,
        italic: bool | None | Unchanged = UNCHANGED,
        underline: Underline | None | Unchanged = UNCHANGED,
        strike: bool | None | Unchanged = UNCHANGED,
        color: str | None | Unchanged = UNCHANGED,
    ) -> WorksheetToolkit:
        """Set the font of cells.

        Parameters
        ----------
        cells : str, optional
            The cells to format, as a range: a block such as 'A1:C3', whole columns
            such as 'B:D', whole rows such as '2:5', or one cell such as 'C3'. Give
            this or ``rows`` and ``columns``, not both.
        rows : list of int, optional
            Row numbers. Defaults to every row in use.
        columns : list of int, optional
            Column numbers. Defaults to every column in use.
        intersections_only : bool, optional
            True, the default, formats the cells where the given rows and the given
            columns cross. False formats every cell in those rows and every cell in
            those columns, which is a cross rather than a block. Only considered when
            both ``rows`` and ``columns`` are given.
        name : str, optional
            Name of the font face, such as 'Calibri'. None reverts the cell back to
            the workbook's default face.
        size : float, optional
            Size in points. None reverts the cell back to the workbook's default size.
        bold : bool, optional
            True draws the text in the bold weight of the face.
        italic : bool, optional
            True draws the text in the italic style of the face.
        underline : str, optional
            'single', 'double', 'singleAccounting' or 'doubleAccounting'. None removes
            the underline.
        strike : bool, optional
            True draws a line through the text.
        color : str, optional
            Hex color code, such as '#f4d2d3'. None clears the color.

        Returns
        -------
        WorksheetToolkit

        Raises
        ------
        TypeError
            If ``rows`` or ``columns`` is given a single number rather than a
            list, or an index that is not an integer.
        ValueError
            If both ``cells`` and ``rows`` or ``columns`` are given, if ``cells``
            cannot be read as a range, or if an index falls outside Excel's grid.

        Examples
        --------
        Bold rows 1 and 2:

        >>> toolkit.set_font(rows=[1, 2], bold=True)

        Size 16 on the four cells where rows 1 and 2 meet columns 1 and 2:

        >>> toolkit.set_font(
        ...     rows=[1, 2], columns=[1, 2], intersections_only=True, size=16
        ... )

        With ``intersections_only=False``, all of rows 1 and 2 and all of columns
        1 and 2, which is a cross rather than a block:

        >>> toolkit.set_font(
        ...     rows=[1, 2], columns=[1, 2], intersections_only=False, color='#000000'
        ... )
        """
        parameters = {
            "name": name,
            "size": size,
            "bold": bold,
            "italic": italic,
            "underline": underline,
            "strike": strike,
            "color": normalize_color(color),
        }

        updates = []
        for cell in iter_cells(self.worksheet, rows, columns, intersections_only, cells):
            # Copy and override, rather than build a new Font from the arguments:
            # a fresh Font would silently reset every attribute this method does
            # not expose, such as vertAlign and scheme.
            new_font = cast(Font, copy(cell.font))
            for attribute, value in parameters.items():
                if value is not UNCHANGED:
                    setattr(new_font, attribute, value)
            updates.append((cell, new_font))

        for cell, font in updates:
            cell.font = font
        return self

    def set_number_format(
        self,
        *,
        number_format: str,
        cells: str | None = None,
        rows: IndexSelection = None,
        columns: IndexSelection = None,
        intersections_only: bool = True,
    ) -> WorksheetToolkit:
        """Set how the value in each cell is displayed.

        The code is Excel's own: '0.00' for two decimal places, '"$"#,##0.00' for
        currency, '0.00%' for a percentage, 'yyyy-mm-dd' for a date. 'General' is
        the default, and ';;;' shows nothing at all while leaving the value in
        place.

        Parameters
        ----------
        number_format : str
            The format code. Required: there is no existing value to leave alone,
            so omitting it cannot mean anything.
        cells : str, optional
            The cells to format, as a range: a block such as 'A1:C3', whole columns
            such as 'B:D', whole rows such as '2:5', or one cell such as 'C3'. Give
            this or ``rows`` and ``columns``, not both.
        rows : list of int, optional
            Row numbers. Defaults to every row in use.
        columns : list of int, optional
            Column numbers. Defaults to every column in use.
        intersections_only : bool, optional
            True, the default, formats the cells where the given rows and the given
            columns cross. False formats every cell in those rows and every cell in
            those columns, which is a cross rather than a block. Only considered when
            both ``rows`` and ``columns`` are given.

        Returns
        -------
        WorksheetToolkit

        Raises
        ------
        TypeError
            If ``number_format`` is not a string, or if ``rows`` or ``columns`` is
            given a single number rather than a list, or an index that is not an
            integer.
        ValueError
            If both ``cells`` and ``rows`` or ``columns`` are given, if ``cells``
            cannot be read as a range, or if an index falls outside Excel's grid.

        Notes
        -----
        Column fitting does not read most of these. ``set_column_best_fit`` renders
        dates and times as Excel displays them and measures everything else as the
        value is stored, so a currency column can come out narrower than it needs
        to be.

        Examples
        --------
        Two decimal places down column B:

        >>> toolkit.set_number_format(number_format='0.00', cells="B:B")

        Currency, where rows 2 and 3 meet column 3:

        >>> toolkit.set_number_format(
        ...     number_format='"$"#,##0.00', rows=[2, 3], columns=[3]
        ... )
        """
        if not isinstance(number_format, str):
            # openpyxl takes the assignment and the workbook then cannot be saved:
            # the stylesheet writer raises on a format code that is not a string.
            raise TypeError(
                f"number_format must be a string, got {number_format!r}. "
                f"Pass 'General' for Excel's default."
            )

        # No two-phase build here, unlike the style setters: there is no object to
        # construct, and assigning a string cannot fail part-way through a range.
        for cell in iter_cells(self.worksheet, rows, columns, intersections_only, cells):
            cell.number_format = number_format
        return self

    def set_zoom_scale(self, zoom_scale: int = 100) -> WorksheetToolkit:
        """Set how far the sheet is zoomed in when it is opened.

        Parameters
        ----------
        zoom_scale : int, optional
            Percentage, 10 to 400. Defaults to 100.

        Returns
        -------
        WorksheetToolkit

        Raises
        ------
        ValueError
            If ``zoom_scale`` is outside 10 to 400.
        """
        if not 10 <= zoom_scale <= 400:
            raise ValueError("zoom_scale must be between 10 and 400")
        self.worksheet.sheet_view.zoomScale = zoom_scale
        return self
