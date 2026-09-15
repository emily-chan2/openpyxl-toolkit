"""Working out which cells a call is about.

Every method that takes a range comes through here, so "A1:C3", whole columns,
whole rows and the four corner numbers all resolve the same way. Index checks
happen before any cell is created: openpyxl will happily materialise a cell
outside Excel's grid, and the workbook can then never be saved.
"""

from numbers import Number

from openpyxl.utils.cell import range_boundaries

from ._limits import MAX_COLUMN, MAX_ROW


def check_indexes(indexes, limit, label):
    """Reject indexes openpyxl would accept but Excel cannot store."""
    for index in indexes:
        # bool is an int, and a float index writes a cell reference like "A1.5"
        # that openpyxl itself can no longer parse when reloading the file.
        if not isinstance(index, int) or isinstance(index, bool):
            raise TypeError(f"{label} index must be an integer, got {index!r}")
        if not 1 <= index <= limit:
            raise ValueError(
                f"{label} {index} is outside the worksheet: {label}s run from 1 to {limit}"
            )


def check_bounds(rows=(), columns=()):
    check_indexes(rows, MAX_ROW, "row")
    check_indexes(columns, MAX_COLUMN, "column")


def resolve_cells(worksheet, cells):
    """Turn a range string into the rows and columns it covers.

    Accepts the three spellings Excel uses:
    - a block such as "A1:C3"
    - whole columns such as "B:D"
    - whole rows such as "2:5"

    An unbounded side is filled in from the worksheet's used range, so "B:B"
    means column B as far as the sheet goes rather than all 1,048,576 rows.
    """
    text = (cells or "").strip()
    if not text:
        raise ValueError("cells needs a range such as 'A1:C3', 'B:B' or '2:5'")

    min_col, min_row, max_col, max_row = range_boundaries(text)
    # "C3:A1" is the same block as "A1:C3"; openpyxl reports it back-to-front.
    if min_row is not None and max_row is not None and min_row > max_row:
        min_row, max_row = max_row, min_row
    if min_col is not None and max_col is not None and min_col > max_col:
        min_col, max_col = max_col, min_col

    rows = list(range(min_row or 1, (max_row if max_row is not None else worksheet.max_row) + 1))
    columns = list(
        range(min_col or 1, (max_col if max_col is not None else worksheet.max_column) + 1)
    )
    return rows, columns


def check_range_arguments(name, cells, start_row, start_column, end_row, end_column):
    """Check the range arguments before they reach openpyxl.

    Raises if both ``cells`` and any of the four coordinates are given, or if
    only some of the coordinates are. openpyxl reports a missing coordinate as
    "expected <class 'int'>" and ignores the coordinates when a range string is
    given as well.
    """
    corners = {
        "start_row": start_row,
        "start_column": start_column,
        "end_row": end_row,
        "end_column": end_column,
    }
    given = [key for key, value in corners.items() if value is not None]

    if cells is not None:
        if given:
            raise ValueError(
                f"give {name} either cells or the four coordinates, not both. "
                f"cells={cells!r} already says which cells to use, "
                f"so {', '.join(given)} would be ignored."
            )
        return

    missing = [key for key in corners if key not in given]
    if missing:
        raise ValueError(
            f"{name} needs either cells, such as 'A1:C3', or all four "
            f"coordinates. Missing: {', '.join(missing)}."
        )


def iter_cells(worksheet, rows=None, columns=None, intersections_only=True, cells=None):
    if cells is not None:
        if rows is not None or columns is not None:
            raise ValueError(
                "give either cells or rows/columns, not both. "
                f"cells={cells!r} already says which cells to use."
            )
        rows, columns = resolve_cells(worksheet, cells)
        # A range names its own block, so there is nothing to intersect or union.
        intersections_only = True

    # Normalize rows and columns
    if rows is None:
        rows = []
    if columns is None:
        columns = []
    if isinstance(rows, Number):
        rows = [rows]
    if isinstance(columns, Number):
        columns = [columns]
    if not rows:
        intersections_only = True
        rows = list(range(1, worksheet.max_row + 1))
    if not columns:
        intersections_only = True
        columns = list(range(1, worksheet.max_column + 1))
    rows = sorted(set(rows))
    columns = sorted(set(columns))
    check_bounds(rows, columns)

    # --- Intersection case ---
    if intersections_only:
        for r in rows:
            for c in columns:
                yield worksheet.cell(row=r, column=c)
        return

    # --- Union case ---
    # The bounds are read once, up front. worksheet.cell() creates a cell that does not
    # exist, so re-reading max_row/max_column inside the loops lets the row sweep
    # grow the sheet and the column sweep then cover rows nobody asked for.
    last_row, last_column = worksheet.max_row, worksheet.max_column
    seen = set()

    # Row sweep (top to bottom, left to right)
    for r in rows:
        for c in range(1, last_column + 1):
            key = (r, c)
            if key not in seen:
                seen.add(key)
                yield worksheet.cell(row=r, column=c)

    # Column sweep (left to right, top to bottom)
    for c in columns:
        for r in range(1, last_row + 1):
            key = (r, c)
            if key not in seen:
                seen.add(key)
                yield worksheet.cell(row=r, column=c)
