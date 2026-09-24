"""Sheet-level settings and the cell-selection mechanics behind every style setter."""

import re
import zipfile

import pytest
from openpyxl import Workbook

from openpyxl_toolkit import WorksheetToolkit


def _grid(worksheet, rows, columns):
    """Give the worksheet a used range of ``rows`` x ``columns`` populated cells."""
    for row in range(1, rows + 1):
        for column in range(1, columns + 1):
            worksheet.cell(row=row, column=column, value=f"r{row}c{column}")
    return worksheet


def _bold_cells(worksheet, through_row=5, through_column=5):
    """Coordinates of every bold cell in the top-left window, as written to the file."""
    return {
        worksheet.cell(row=row, column=column).coordinate
        for row in range(1, through_row + 1)
        for column in range(1, through_column + 1)
        if worksheet.cell(row=row, column=column).font.bold
    }


# --- sheet-level settings ----------------------------------------------------


def test_freeze_panes_at_b2_survives_the_round_trip(sheet, roundtrip):
    WorksheetToolkit(sheet).freeze_panes("B2")

    assert roundtrip(sheet).freeze_panes == "B2"


def test_freeze_panes_accepts_a_lower_case_reference(sheet, roundtrip):
    WorksheetToolkit(sheet).freeze_panes("b2")

    assert roundtrip(sheet).freeze_panes == "B2"


def test_freezing_at_a1_is_the_same_as_passing_none(sheet, tmp_path):
    """A1 names a scrolling area starting at the corner, so it freezes nothing.

    Compared as saved files rather than as attributes. Unfreezing has to clear the
    pane and the three split selections together, and openpyxl clears only the
    pane, which is what makes some versions of Excel offer to repair the file.
    """

    def saved(unfreeze_with):
        workbook = Workbook()
        worksheet = workbook.active
        worksheet["A1"] = "x"
        toolkit = WorksheetToolkit(worksheet)
        toolkit.freeze_panes("B2")
        toolkit.freeze_panes(unfreeze_with)
        path = tmp_path / f"{unfreeze_with}.xlsx"
        workbook.save(path)
        with zipfile.ZipFile(path) as archive:
            return archive.read("xl/worksheets/sheet1.xml")

    assert saved("A1") == saved(None)
    assert saved("a1") == saved(None)


def test_an_unfrozen_sheet_is_indistinguishable_from_one_never_frozen(sheet, tmp_path):
    """The property Excel actually reacts to, asserted on the saved file.

    Removing a pane through openpyxl alone leaves three selection elements, two
    of them naming panes that no longer exist. Opening such a file in Excel
    produces the offer to repair it; the file this produces opens cleanly. Both
    were checked by hand once, and this keeps the difference from coming back.
    """

    def saved(name, build):
        workbook = Workbook()
        worksheet = workbook.active
        worksheet["A1"] = "x"
        build(WorksheetToolkit(worksheet), worksheet)
        path = tmp_path / f"{name}.xlsx"
        workbook.save(path)
        with zipfile.ZipFile(path) as archive:
            return archive.read("xl/worksheets/sheet1.xml")

    never_frozen = saved("never", lambda toolkit, worksheet: None)
    unfrozen = saved(
        "unfrozen",
        lambda toolkit, worksheet: (toolkit.freeze_panes("B2"), toolkit.freeze_panes(None)),
    )
    left_behind = saved(
        "raw",
        lambda toolkit, worksheet: (
            toolkit.freeze_panes("B2"),
            setattr(worksheet, "freeze_panes", None),
        ),
    )

    assert unfrozen == never_frozen
    assert left_behind != never_frozen


def test_a_sheet_that_was_never_frozen_has_one_selection(sheet):
    """What unfreezing has to get back to, and what the comparison above rests on."""
    assert len(sheet.sheet_view.selection) == 1


@pytest.mark.parametrize("cell", ["", "nonsense", "B", "2", "B2:C3", " B2 "])
def test_freeze_panes_rejects_anything_that_is_not_one_cell(sheet, cell):
    """An empty string is not a spelling of None: openpyxl took it and unfroze
    without clearing the selections."""
    with pytest.raises(ValueError, match="one cell reference"):
        WorksheetToolkit(sheet).freeze_panes(cell)


def test_freeze_panes_rejects_a_cell_outside_the_grid(sheet):
    with pytest.raises(ValueError, match="outside the worksheet"):
        WorksheetToolkit(sheet).freeze_panes("ZZZ99999999")


def test_freeze_panes_needs_an_argument(sheet):
    with pytest.raises(TypeError, match="cell"):
        WorksheetToolkit(sheet).freeze_panes()


def test_merged_range_survives_the_round_trip(sheet, roundtrip):
    WorksheetToolkit(sheet).merge_cells(cells="A1:C2")

    assert [str(cell_range) for cell_range in roundtrip(sheet).merged_cells.ranges] == ["A1:C2"]


def test_merging_by_coordinates_produces_the_same_range_as_cells(sheet, roundtrip):
    WorksheetToolkit(sheet).merge_cells(start_row=1, start_column=1, end_row=2, end_column=3)

    assert [str(cell_range) for cell_range in roundtrip(sheet).merged_cells.ranges] == ["A1:C2"]


def test_zoom_scale_survives_the_round_trip(sheet, roundtrip):
    WorksheetToolkit(sheet).set_zoom_scale(zoom_scale=85)

    assert roundtrip(sheet).sheet_view.zoomScale == 85


@pytest.mark.parametrize("zoom_scale", [10, 400])
def test_zoom_scale_bounds_are_inclusive(sheet, roundtrip, zoom_scale):
    WorksheetToolkit(sheet).set_zoom_scale(zoom_scale=zoom_scale)

    assert roundtrip(sheet).sheet_view.zoomScale == zoom_scale


@pytest.mark.parametrize("zoom_scale", [9, 401, 0, -100])
def test_zoom_scale_outside_the_allowed_span_is_rejected(sheet, zoom_scale):
    with pytest.raises(ValueError):
        WorksheetToolkit(sheet).set_zoom_scale(zoom_scale=zoom_scale)


def test_tab_color_survives_the_round_trip(sheet, roundtrip):
    WorksheetToolkit(sheet).set_tab_color("#1d3557")

    assert roundtrip(sheet).sheet_properties.tabColor.rgb == "FF1D3557"


def test_a_six_digit_tab_color_is_stored_opaque(sheet, roundtrip):
    """openpyxl pads a six digit hex with alpha 00, which is fully transparent."""
    WorksheetToolkit(sheet).set_tab_color("1d3557")

    assert roundtrip(sheet).sheet_properties.tabColor.rgb == "FF1D3557"


def test_a_tab_color_matches_the_argb_the_other_methods_write(sheet, roundtrip):
    """One hex handed to two methods has to reach the file as one color."""
    toolkit = WorksheetToolkit(sheet)
    toolkit.set_tab_color("#ff0000")
    toolkit.set_font(cells="A1", color="#ff0000")

    reloaded = roundtrip(sheet)
    assert reloaded.sheet_properties.tabColor.rgb == reloaded["A1"].font.color.rgb


def test_tab_color_none_clears_it(sheet, roundtrip):
    toolkit = WorksheetToolkit(sheet)
    toolkit.set_tab_color("#1d3557")
    toolkit.set_tab_color(None)

    assert roundtrip(sheet).sheet_properties.tabColor is None


def test_a_sheet_with_no_tab_color_set_has_none(sheet, roundtrip):
    """The clearing test would pass without reading anything if this were not true."""
    assert roundtrip(sheet).sheet_properties.tabColor is None


def test_hiding_gridlines_survives_the_round_trip(sheet, roundtrip):
    WorksheetToolkit(sheet).set_gridline_visibility(visible=False)

    assert roundtrip(sheet).sheet_view.showGridLines is False


def test_showing_gridlines_again_survives_the_round_trip(sheet, roundtrip):
    toolkit = WorksheetToolkit(sheet)
    toolkit.set_gridline_visibility(visible=False)
    toolkit.set_gridline_visibility(visible=True)

    assert roundtrip(sheet).sheet_view.showGridLines is True


def test_a_new_sheet_leaves_gridlines_alone(sheet, roundtrip):
    """Untouched is not the same as shown: openpyxl leaves the setting unwritten."""
    assert roundtrip(sheet).sheet_view.showGridLines is None


def test_hiding_gridlines_leaves_printed_gridlines_alone(sheet, roundtrip):
    """Excel keeps screen and print gridlines apart, and so does this."""
    WorksheetToolkit(sheet).set_gridline_visibility(visible=False)

    assert roundtrip(sheet).print_options.gridLines is None


def test_hiding_gridlines_leaves_cell_borders_drawn(sheet, roundtrip):
    """The point of hiding the grid is that deliberate borders still show."""
    toolkit = WorksheetToolkit(sheet)
    toolkit.set_border(cells="A1", sides=("bottom",), style="thin")
    toolkit.set_gridline_visibility(visible=False)

    reloaded = roundtrip(sheet)
    assert reloaded.sheet_view.showGridLines is False
    assert reloaded["A1"].border.bottom.style == "thin"


def test_autofilter_survives_the_round_trip(sheet, roundtrip):
    _grid(sheet, 4, 3)

    WorksheetToolkit(sheet).set_autofilter(cells="A1:C4")

    assert roundtrip(sheet).auto_filter.ref == "A1:C4"


def test_autofilter_by_coordinates_matches_the_range_string(sheet, roundtrip):
    _grid(sheet, 4, 3)

    WorksheetToolkit(sheet).set_autofilter(start_row=1, start_column=1, end_row=4, end_column=3)

    assert roundtrip(sheet).auto_filter.ref == "A1:C4"


def test_autofilter_none_removes_it(sheet, roundtrip):
    _grid(sheet, 4, 3)
    toolkit = WorksheetToolkit(sheet)
    toolkit.set_autofilter(cells="A1:C4")

    toolkit.set_autofilter(cells=None)

    assert roundtrip(sheet).auto_filter.ref is None


def test_a_sheet_with_no_autofilter_has_none(sheet, roundtrip):
    """Without this the removal test would pass by reading nothing."""
    assert roundtrip(sheet).auto_filter.ref is None


def test_an_autofilter_carries_no_criteria(sheet, roundtrip):
    """The control, not a rule. openpyxl can write a rule but hides no rows."""
    _grid(sheet, 4, 3)

    WorksheetToolkit(sheet).set_autofilter(cells="A1:C4")

    reloaded = roundtrip(sheet)
    assert reloaded.auto_filter.filterColumn == []
    assert [reloaded.row_dimensions[r].hidden for r in range(1, 5)] == [False] * 4


def test_autofilter_needs_a_range(sheet):
    with pytest.raises(ValueError, match="set_autofilter needs either"):
        WorksheetToolkit(sheet).set_autofilter()


def test_autofilter_rejects_both_spellings_at_once(sheet):
    with pytest.raises(ValueError, match="not both"):
        WorksheetToolkit(sheet).set_autofilter(cells="A1:C4", start_row=1)


@pytest.mark.parametrize("cells", ["A:C", "B:B"])
def test_autofilter_accepts_whole_columns(sheet, roundtrip, cells):
    """Excel works out the extent on opening, so the filter covers the data
    however much of it there turns out to be. Checked by hand in Excel once: the
    arrows appear and the filter works."""
    _grid(sheet, 4, 3)

    WorksheetToolkit(sheet).set_autofilter(cells=cells)

    assert roundtrip(sheet).auto_filter.ref == cells


@pytest.mark.parametrize("cells", ["1:5", "2:2"])
def test_autofilter_rejects_whole_rows(sheet, cells):
    """Not a form Excel has. openpyxl refuses it too, with its raw pattern."""
    _grid(sheet, 4, 3)

    with pytest.raises(ValueError, match="whole columns"):
        WorksheetToolkit(sheet).set_autofilter(cells=cells)


def test_autofilter_can_be_set_before_the_data_is_written(sheet, roundtrip):
    """A concrete range does not depend on what is on the sheet yet."""
    WorksheetToolkit(sheet).set_autofilter(cells="A1:C4")

    for row in range(1, 5):
        for column in range(1, 4):
            sheet.cell(row=row, column=column, value="x")

    assert roundtrip(sheet).auto_filter.ref == "A1:C4"


def test_autofilter_accepts_a_lower_case_range(sheet, roundtrip):
    _grid(sheet, 4, 3)

    WorksheetToolkit(sheet).set_autofilter(cells="a1:c4")

    assert roundtrip(sheet).auto_filter.ref == "A1:C4"


def test_autofilter_outside_the_grid_is_rejected(sheet):
    with pytest.raises(ValueError, match="outside the worksheet"):
        WorksheetToolkit(sheet).set_autofilter(
            start_row=1, start_column=1, end_row=4, end_column=20000
        )


def test_hiding_a_sheet_survives_the_round_trip(sheet, roundtrip):
    second = sheet.parent.create_sheet("Second")

    WorksheetToolkit(second).set_sheet_visibility(state="hidden")

    assert roundtrip(sheet).parent["Second"].sheet_state == "hidden"


def test_very_hidden_is_written_with_excels_spelling(sheet, roundtrip):
    """The package says very_hidden; the file has to say veryHidden."""
    second = sheet.parent.create_sheet("Second")

    WorksheetToolkit(second).set_sheet_visibility(state="very_hidden")

    assert roundtrip(sheet).parent["Second"].sheet_state == "veryHidden"


def test_a_hidden_sheet_can_be_shown_again(sheet, roundtrip):
    second = sheet.parent.create_sheet("Second")
    toolkit = WorksheetToolkit(second)
    toolkit.set_sheet_visibility(state="very_hidden")

    toolkit.set_sheet_visibility(state="visible")

    assert roundtrip(sheet).parent["Second"].sheet_state == "visible"


def test_excels_own_spelling_is_rejected_and_the_message_gives_the_right_one(sheet):
    second = sheet.parent.create_sheet("Second")

    with pytest.raises(ValueError, match="very_hidden"):
        WorksheetToolkit(second).set_sheet_visibility(state="veryHidden")


def test_an_unknown_sheet_state_is_rejected(sheet):
    second = sheet.parent.create_sheet("Second")

    with pytest.raises(ValueError, match="state must be one of"):
        WorksheetToolkit(second).set_sheet_visibility(state="nope")


def test_hiding_every_sheet_is_allowed_while_the_code_runs(sheet):
    """A workbook part way through being rearranged is not a file yet.

    Refusing here would mean the new sheet had to be shown before the old one
    could be hidden, which is an ordering the caller should not have to know.
    """
    second = sheet.parent.create_sheet("Second")

    WorksheetToolkit(second).set_sheet_visibility(state="hidden")
    WorksheetToolkit(sheet).set_sheet_visibility(state="hidden")

    assert [s.sheet_state for s in sheet.parent.worksheets] == ["hidden", "hidden"]


def test_saving_a_workbook_with_nothing_visible_still_fails(sheet, tmp_path):
    """Allowed in memory, refused in a file. openpyxl is what refuses it."""
    WorksheetToolkit(sheet).set_sheet_visibility(state="hidden")

    with pytest.raises((ValueError, IndexError)):
        sheet.parent.save(tmp_path / "book.xlsx")


def test_a_workbook_whose_sheets_are_all_visible_still_saves(sheet, roundtrip):
    """The guard would pass by rejecting everything if this were not checked."""
    second = sheet.parent.create_sheet("Second")
    WorksheetToolkit(second).set_sheet_visibility(state="hidden")

    assert roundtrip(sheet).parent.sheetnames == ["Sheet", "Second"]


def test_zoom_scale_needs_a_value(sheet):
    with pytest.raises(TypeError, match="zoom_scale"):
        WorksheetToolkit(sheet).set_zoom_scale()


def test_sheet_level_methods_return_the_toolkit_for_chaining(sheet):
    toolkit = WorksheetToolkit(sheet)

    chained = toolkit.freeze_panes("B2").merge_cells(cells="A1:B1").set_zoom_scale(zoom_scale=120)

    assert chained is toolkit


# --- cell selection, observed through set_font -------------------------------


def test_selecting_rows_only_touches_every_column_of_those_rows(sheet, roundtrip):
    _grid(sheet, 3, 3)

    WorksheetToolkit(sheet).set_font(rows=[2], bold=True)

    assert _bold_cells(roundtrip(sheet)) == {"A2", "B2", "C2"}


def test_selecting_columns_only_touches_every_row_of_those_columns(sheet, roundtrip):
    _grid(sheet, 3, 3)

    WorksheetToolkit(sheet).set_font(columns=[2], bold=True)

    assert _bold_cells(roundtrip(sheet)) == {"B1", "B2", "B3"}


def test_selecting_neither_rows_nor_columns_touches_the_whole_used_range(sheet, roundtrip):
    _grid(sheet, 2, 2)

    WorksheetToolkit(sheet).set_font(bold=True)

    assert _bold_cells(roundtrip(sheet)) == {"A1", "A2", "B1", "B2"}


def test_intersection_selection_touches_only_the_crossings(sheet, roundtrip):
    _grid(sheet, 3, 3)

    WorksheetToolkit(sheet).set_font(
        rows=[1, 2], columns=[1, 3], intersections_only=True, bold=True
    )

    assert _bold_cells(roundtrip(sheet)) == {"A1", "A2", "C1", "C2"}


def test_union_selection_touches_whole_rows_and_whole_columns(sheet, roundtrip):
    _grid(sheet, 3, 3)

    WorksheetToolkit(sheet).set_font(rows=[1], columns=[1], intersections_only=False, bold=True)

    assert _bold_cells(roundtrip(sheet)) == {"A1", "B1", "C1", "A2", "A3"}


def test_duplicate_and_unordered_selectors_touch_each_cell_once(sheet, roundtrip):
    _grid(sheet, 3, 3)

    WorksheetToolkit(sheet).set_font(rows=[3, 1, 1], columns=[2, 2], bold=True)

    assert _bold_cells(roundtrip(sheet)) == {"B1", "B3"}


def test_a_later_font_attribute_does_not_reset_an_earlier_one(sheet, roundtrip):
    """Overlapping selections merge: the second call keeps the first call's bold."""
    _grid(sheet, 2, 3)
    toolkit = WorksheetToolkit(sheet)

    toolkit.set_font(rows=[1, 2], columns=[1, 2], bold=True)
    toolkit.set_font(rows=[1], columns=[1, 2, 3], size=16)

    reloaded = roundtrip(sheet)
    assert [
        (reloaded.cell(row=1, column=c).font.bold, reloaded.cell(row=1, column=c).font.size)
        for c in (1, 2, 3)
    ] == [(True, 16.0), (True, 16.0), (False, 16.0)]


def test_a_later_selection_does_not_reset_cells_outside_it(sheet, roundtrip):
    """Cells the second call never selects keep everything the first call gave them."""
    _grid(sheet, 2, 3)
    toolkit = WorksheetToolkit(sheet)

    toolkit.set_font(rows=[1, 2], columns=[1, 2], bold=True)
    toolkit.set_font(rows=[1], columns=[1, 2, 3], size=16)

    reloaded = roundtrip(sheet)
    assert [
        (reloaded.cell(row=2, column=c).font.bold, reloaded.cell(row=2, column=c).font.size)
        for c in (1, 2)
    ] == [(True, 11.0), (True, 11.0)]


# --- known defects -----------------------------------------------------------


def test_naming_a_row_outside_the_data_styles_it(sheet, roundtrip):
    """Cells named explicitly are styled even where the sheet holds no data yet.

    This does extend the used range, and that is the intended trade: the only way
    to grow the sheet is to ask for it by name. Skipping those cells instead would
    make the call quietly do less than it was asked to.
    """
    sheet["A1"] = "only cell"

    WorksheetToolkit(sheet).set_font(rows=[1, 2], columns=[1, 2], bold=True)

    assert _bold_cells(roundtrip(sheet)) == {"A1", "A2", "B1", "B2"}


def test_the_default_selection_does_not_extend_the_sheet(sheet):
    """rows=None means the used range, and reading it must not grow it.

    set_column_best_fit and the selection defaults both read the used range, so an
    automatic call that quietly enlarged it would compound on every later call.
    """
    sheet["A1"] = "only cell"

    WorksheetToolkit(sheet).set_font(bold=True)

    assert (sheet.max_row, sheet.max_column) == (1, 1)


def test_union_selection_does_not_reach_rows_the_caller_never_asked_for(sheet, roundtrip):
    """Row 3 is in neither the requested rows nor the data, so it stays plain."""
    _grid(sheet, 2, 2)

    WorksheetToolkit(sheet).set_font(rows=[4], columns=[1], intersections_only=False, bold=True)

    assert _bold_cells(roundtrip(sheet)) == {"A1", "A2", "A4", "B4"}


def test_a_fill_across_a_merged_range_reaches_every_cell_in_the_file(sheet, tmp_path):
    """Checked against the saved file, not a reload.

    openpyxl replaces the non-anchor cells of a merge with MergedCell objects when it
    reads a workbook, and drops their style while doing so. The style is in the file:
    reading it back is what loses it, so a round-trip assertion would be testing
    openpyxl rather than this library.
    """
    toolkit = WorksheetToolkit(sheet)
    sheet["A1"] = "heading"
    toolkit.merge_cells(cells="A1:C1")

    toolkit.set_fill(rows=[1], columns=[1, 2, 3], fill_type="solid", start_color="#FFD966")

    path = tmp_path / "book.xlsx"
    sheet.parent.save(path)
    with zipfile.ZipFile(path) as archive:
        worksheet = archive.read("xl/worksheets/sheet1.xml").decode()
        styles = archive.read("xl/styles.xml").decode()

    row = re.search(r'<row r="1".*?</row>', worksheet, re.S).group()
    used = re.findall(r'<c r="([A-C]1)" s="(\d+)"', row)
    assert [coordinate for coordinate, _ in used] == ["A1", "B1", "C1"]
    assert len({style for _, style in used}) == 1, "the three cells disagree on style"

    # cellStyleXfs precedes cellXfs, so the cell's style index counts within cellXfs
    cell_formats = re.search(r"<cellXfs.*?</cellXfs>", styles, re.S).group()
    used_format = re.findall(r"<xf [^>]*/>", cell_formats)[int(used[0][1])]
    fill_index = int(re.search(r'fillId="(\d+)"', used_format).group(1))
    assert "FFFFD966" in re.findall(r"<fill>.*?</fill>", styles, re.S)[fill_index]


def test_unfreezing_leaves_no_pane_state_behind(sheet, roundtrip):
    toolkit = WorksheetToolkit(sheet)
    toolkit.freeze_panes("B2")

    toolkit.freeze_panes(None)

    reloaded = roundtrip(sheet)
    assert not reloaded.freeze_panes
    assert [selection.pane for selection in reloaded.sheet_view.selection] == [None]


def test_a_column_beyond_the_grid_is_rejected(sheet):
    """openpyxl creates such a cell unvalidated, and the workbook cannot be saved."""
    with pytest.raises(ValueError, match="column 20000"):
        WorksheetToolkit(sheet).set_font(rows=[1], columns=[20000], bold=True)


def test_a_row_beyond_the_grid_is_rejected(sheet):
    with pytest.raises(ValueError, match="row 2000000"):
        WorksheetToolkit(sheet).set_font(rows=[2000000], columns=[1], bold=True)


def test_a_column_of_zero_is_rejected(sheet):
    with pytest.raises(ValueError, match="column 0"):
        WorksheetToolkit(sheet).set_font(rows=[1], columns=[0], bold=True)


def test_rejecting_a_bad_index_leaves_the_workbook_saveable(sheet, roundtrip):
    """Rejecting early matters: a materialised out-of-range cell cannot be undone."""
    _grid(sheet, 2, 2)
    toolkit = WorksheetToolkit(sheet)

    with pytest.raises(ValueError):
        toolkit.set_font(rows=[1], columns=[20000], bold=True)

    assert roundtrip(sheet)["A1"].value == "r1c1"


def test_the_dimension_setters_are_bounds_checked_too(sheet):
    """16385 is past Excel's limit but inside get_column_letter's range."""
    toolkit = WorksheetToolkit(sheet)
    for call in (
        lambda: toolkit.set_column_width(columns=[16385], width=12),
        lambda: toolkit.set_column_best_fit(columns=[16385]),
    ):
        with pytest.raises(ValueError, match="outside the worksheet"):
            call()


def test_set_row_height_is_bounds_checked(sheet):
    with pytest.raises(ValueError, match="outside the worksheet"):
        WorksheetToolkit(sheet).set_row_height(rows=[2000000], height=20)


def test_best_fit_rejects_a_bad_column_before_materialising_it(sheet, roundtrip):
    """get_column_letter raises only after ws.cell has made the sheet unsaveable."""
    _grid(sheet, 2, 2)
    with pytest.raises(ValueError):
        WorksheetToolkit(sheet).set_column_best_fit(columns=[1, 20000])

    assert roundtrip(sheet)["A1"].value == "r1c1"


def test_a_fractional_index_is_rejected(sheet):
    """A float index writes a cell reference like A1.5 that openpyxl cannot reload."""
    toolkit = WorksheetToolkit(sheet)
    with pytest.raises(TypeError, match="must be an integer"):
        toolkit.set_font(rows=[1.5], columns=[1], bold=True)
    with pytest.raises(TypeError, match="must be an integer"):
        toolkit.set_row_height(rows=[1.5], height=20)


def test_the_toolkit_rejects_anything_that_is_not_a_worksheet(sheet):
    """Passing the workbook is the easy slip, and it used to fail much later."""
    with pytest.raises(TypeError, match="expected a Worksheet, got Workbook"):
        WorksheetToolkit(sheet.parent)
    with pytest.raises(TypeError, match="got str"):
        WorksheetToolkit("Sheet1")


def test_unmerge_cells_undoes_a_merge(sheet, roundtrip):
    toolkit = WorksheetToolkit(sheet)
    sheet["A1"] = "heading"
    toolkit.merge_cells(cells="A1:C1")

    toolkit.unmerge_cells(cells="A1:C1")

    assert list(roundtrip(sheet).merged_cells.ranges) == []


def test_unmerge_cells_accepts_coordinates_too(sheet, roundtrip):
    toolkit = WorksheetToolkit(sheet)
    sheet["A1"] = "heading"
    toolkit.merge_cells(start_row=1, start_column=1, end_row=1, end_column=3)

    toolkit.unmerge_cells(start_row=1, start_column=1, end_row=1, end_column=3)

    assert list(roundtrip(sheet).merged_cells.ranges) == []


def test_unmerging_a_range_that_is_not_merged_does_nothing(sheet):
    """Unmerging a range that is already unmerged is not an error."""
    WorksheetToolkit(sheet).unmerge_cells(cells="A5:C5")


def test_unmerging_is_repeatable(sheet, roundtrip):
    toolkit = WorksheetToolkit(sheet)
    sheet["A1"] = "heading"
    toolkit.merge_cells(cells="A1:C1")

    toolkit.unmerge_cells(cells="A1:C1").unmerge_cells(cells="A1:C1")

    assert list(roundtrip(sheet).merged_cells.ranges) == []


def test_an_unparseable_range_still_raises(sheet):
    """Only the not-merged case is forgiven, not a typo that is not a range at all."""
    with pytest.raises(ValueError):
        WorksheetToolkit(sheet).unmerge_cells(cells="nonsense")


@pytest.mark.parametrize("method", ["merge_cells", "unmerge_cells"])
def test_a_range_with_no_arguments_says_what_is_missing(sheet, method):
    """openpyxl's own message for this is 'expected <class int>'."""
    with pytest.raises(ValueError, match="start_row, start_column, end_row, end_column"):
        getattr(WorksheetToolkit(sheet), method)()


def test_a_partial_range_says_which_corner_is_missing(sheet):
    with pytest.raises(ValueError, match="end_column"):
        WorksheetToolkit(sheet).merge_cells(start_row=1, start_column=1, end_row=1)


def test_the_unchanged_sentinel_reads_as_words(sheet):
    """It appears in every signature help() and an editor tooltip render."""
    import inspect

    from openpyxl_toolkit._sentinel import UNCHANGED

    assert repr(UNCHANGED) == "<unchanged>"
    assert "<unchanged>" in str(inspect.signature(WorksheetToolkit.set_font))
    assert "object object at" not in str(inspect.signature(WorksheetToolkit.set_font))


def test_zoom_scale_takes_its_value_positionally_or_by_keyword(sheet, roundtrip):
    """One argument the method name already explains, so the keyword is optional."""
    WorksheetToolkit(sheet).set_zoom_scale(85)
    assert roundtrip(sheet).sheet_view.zoomScale == 85

    WorksheetToolkit(sheet).set_zoom_scale(zoom_scale=140)
    assert roundtrip(sheet).sheet_view.zoomScale == 140


@pytest.mark.parametrize(
    "call",
    [
        lambda t: t.set_font(rows=1, bold=True),
        lambda t: t.set_font(columns=1, bold=True),
        lambda t: t.set_fill(rows=1, fill_type="solid", start_color="#ff0000"),
        lambda t: t.set_alignment(columns=1, horizontal="center"),
        lambda t: t.set_border(rows=1, style="thin"),
        lambda t: t.set_column_width(columns=1, width=10),
        lambda t: t.set_row_height(rows=1, height=10),
        lambda t: t.set_column_best_fit(columns=1),
    ],
    ids=lambda call: "",
)
def test_a_single_row_or_column_number_is_rejected(sheet, call):
    """rows= and columns= take a list. A bare number used to be wrapped in one."""
    with pytest.raises(TypeError, match=r"takes a list of numbers"):
        call(WorksheetToolkit(sheet))


def test_the_rejection_shows_the_list_to_write_instead(sheet):
    with pytest.raises(TypeError, match=r"Write columns=\[3\] rather than columns=3"):
        WorksheetToolkit(sheet).set_column_width(columns=3, width=10)
