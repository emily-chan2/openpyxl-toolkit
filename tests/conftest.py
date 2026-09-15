"""Shared fixtures.

Every behavioural test in this suite goes through ``roundtrip``: the only failure
mode that matters for this library is a style that looks right in memory and does
not survive being written to a real .xlsx file.
"""

import pytest
from openpyxl import Workbook, load_workbook

# A note that has cost three tests so far: cell.font, cell.fill, cell.border and
# cell.alignment return a StyleProxy, not the style object itself. Two proxies never
# compare equal to each other, and Proxy.__class__() is not a usable constructor.
# Compare individual attributes, or unwrap with copy() first.


@pytest.fixture
def sheet():
    """A fresh, empty worksheet."""
    return Workbook().active


@pytest.fixture
def roundtrip(tmp_path):
    """Save the worksheet's workbook, reload it, and return the same sheet."""

    def _roundtrip(ws):
        path = tmp_path / "book.xlsx"
        ws.parent.save(path)
        return load_workbook(path)[ws.title]

    return _roundtrip
