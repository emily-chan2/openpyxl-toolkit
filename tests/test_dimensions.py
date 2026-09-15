"""Column widths, row heights and best-fit sizing, as observed in the saved file."""

from datetime import datetime

import pytest
from openpyxl.styles import Font

from openpyxl_toolkit import WorksheetToolkit

EXCEL_MAX_COLUMN_WIDTH = 255


# --------------------------------------------------------------------------
# Characterisation: behaviour that is currently correct and must stay correct
# --------------------------------------------------------------------------


def test_column_width_survives_the_round_trip(sheet, roundtrip):
    WorksheetToolkit(sheet).set_column_width(columns=[1], width=21.5)

    assert roundtrip(sheet).column_dimensions["A"].width == 21.5


def test_row_height_survives_the_round_trip(sheet, roundtrip):
    WorksheetToolkit(sheet).set_row_height(rows=[1], height=44.5)

    assert roundtrip(sheet).row_dimensions[1].height == 44.5


def test_a_single_column_number_is_treated_as_a_one_column_list(sheet, roundtrip):
    toolkit = WorksheetToolkit(sheet)
    toolkit.set_column_width(columns=1, width=21)
    toolkit.set_column_width(columns=[2], width=21)

    ws = roundtrip(sheet)
    assert ws.column_dimensions["A"].width == ws.column_dimensions["B"].width


def test_a_single_row_number_is_treated_as_a_one_row_list(sheet, roundtrip):
    toolkit = WorksheetToolkit(sheet)
    toolkit.set_row_height(rows=1, height=33)
    toolkit.set_row_height(rows=[2], height=33)

    ws = roundtrip(sheet)
    assert ws.row_dimensions[1].height == ws.row_dimensions[2].height


def test_width_with_no_columns_applies_to_every_used_column(sheet, roundtrip):
    sheet["A1"] = "a"
    sheet["B1"] = "b"
    sheet["C1"] = "c"

    WorksheetToolkit(sheet).set_column_width(width=17)

    ws = roundtrip(sheet)
    assert [ws.column_dimensions[letter].width for letter in "ABC"] == [17, 17, 17]


def test_height_with_no_rows_applies_to_every_used_row(sheet, roundtrip):
    sheet["A1"] = "a"
    sheet["A2"] = "b"
    sheet["A3"] = "c"

    WorksheetToolkit(sheet).set_row_height(height=27)

    ws = roundtrip(sheet)
    assert [ws.row_dimensions[row].height for row in (1, 2, 3)] == [27, 27, 27]


def test_best_fit_makes_a_long_column_wider_than_a_short_one(sheet, roundtrip):
    sheet["A1"] = "a very long piece of text indeed"
    sheet["B1"] = "hi"

    WorksheetToolkit(sheet).set_column_best_fit()

    ws = roundtrip(sheet)
    assert ws.column_dimensions["A"].width > ws.column_dimensions["B"].width


def test_best_fit_skips_rows_listed_in_ignore_rows(sheet, roundtrip):
    """Column A's oversized header is ignored, so it fits like control column B."""
    sheet["A1"] = "H" * 40
    sheet["A2"] = "short"
    sheet["B2"] = "short"

    WorksheetToolkit(sheet).set_column_best_fit(columns=[1, 2], ignore_rows=[1])

    ws = roundtrip(sheet)
    assert ws.column_dimensions["A"].width == ws.column_dimensions["B"].width


def test_best_fit_skips_formula_cells_by_default(sheet, roundtrip):
    """Column A is sized by its literal cell only, like control column B."""
    sheet["A1"] = '=REPT("x",100)'
    sheet["A2"] = "abc"
    sheet["B2"] = "abc"

    WorksheetToolkit(sheet).set_column_best_fit(columns=[1, 2])

    ws = roundtrip(sheet)
    assert ws.column_dimensions["A"].width == ws.column_dimensions["B"].width


def test_best_fit_measures_formula_text_when_ignore_formulas_is_false(sheet, roundtrip):
    sheet["A1"] = '=REPT("x",100)'
    sheet["A2"] = "abc"
    sheet["B2"] = "abc"

    WorksheetToolkit(sheet).set_column_best_fit(columns=[1, 2], ignore_formulas=False)

    ws = roundtrip(sheet)
    assert ws.column_dimensions["A"].width > ws.column_dimensions["B"].width


def test_best_fit_adds_padding_on_top_of_the_measured_width(sheet, roundtrip):
    sheet["A1"] = "abcde"
    sheet["B1"] = "abcde"

    toolkit = WorksheetToolkit(sheet)
    toolkit.set_column_best_fit(columns=1, padding=0)
    toolkit.set_column_best_fit(columns=2, padding=10)

    ws = roundtrip(sheet)
    difference = ws.column_dimensions["B"].width - ws.column_dimensions["A"].width
    assert difference == pytest.approx(10)


def test_setting_a_row_height_leaves_an_earlier_column_width_intact(sheet, roundtrip):
    """The library's core promise: a later call adds to the layout, it does not replace it."""
    toolkit = WorksheetToolkit(sheet)
    toolkit.set_column_width(columns=1, width=25)
    toolkit.set_row_height(rows=1, height=40)

    ws = roundtrip(sheet)
    assert ws.column_dimensions["A"].width == 25


def test_setting_one_column_width_leaves_another_column_width_intact(sheet, roundtrip):
    toolkit = WorksheetToolkit(sheet)
    toolkit.set_column_width(columns=1, width=25)
    toolkit.set_column_width(columns=2, width=9)

    ws = roundtrip(sheet)
    assert ws.column_dimensions["A"].width == 25


def test_best_fit_leaves_an_explicit_width_on_an_untouched_column_intact(sheet, roundtrip):
    sheet["A1"] = "x"
    sheet["B1"] = "a much longer value"

    toolkit = WorksheetToolkit(sheet)
    toolkit.set_column_width(columns=1, width=42)
    toolkit.set_column_best_fit(columns=2)

    ws = roundtrip(sheet)
    assert ws.column_dimensions["A"].width == 42


def test_dimension_setters_return_the_toolkit_for_chaining(sheet):
    toolkit = WorksheetToolkit(sheet)

    assert toolkit.set_column_width(columns=1, width=10) is toolkit
    assert toolkit.set_row_height(rows=1, height=10) is toolkit
    assert toolkit.set_column_best_fit(columns=1) is toolkit


# --------------------------------------------------------------------------
# Known defects: each test is what should pass once the bug is fixed
# --------------------------------------------------------------------------


def test_column_width_applies_to_non_anchor_merged_columns(sheet, roundtrip):
    toolkit = WorksheetToolkit(sheet)
    toolkit.merge_cells(cells="A1:D1")
    toolkit.set_column_width(columns=[1, 2, 3], width=15)

    ws = roundtrip(sheet)
    assert [ws.column_dimensions[letter].width for letter in "ABC"] == [15, 15, 15]


def test_best_fit_applies_to_non_anchor_merged_columns(sheet, roundtrip):
    sheet["A2"] = "a" * 20
    sheet["B2"] = "b" * 5
    sheet["C2"] = "c" * 10

    toolkit = WorksheetToolkit(sheet)
    toolkit.merge_cells(cells="A1:D1")
    toolkit.set_column_best_fit(columns=[1, 2, 3])

    ws = roundtrip(sheet)
    widths = [ws.column_dimensions[letter].width for letter in "ABC"]
    assert widths[0] > widths[2] > widths[1]


def test_best_fit_handles_a_font_with_no_explicit_size(sheet, roundtrip):
    """A font that inherits its size should fit like the same text at the default size."""
    sheet["A1"] = "hello world"
    sheet["A1"].font = Font(bold=True)
    sheet["B1"] = "hello world"

    WorksheetToolkit(sheet).set_column_best_fit(columns=[1, 2])

    ws = roundtrip(sheet)
    assert ws.column_dimensions["A"].width == pytest.approx(
        ws.column_dimensions["B"].width, rel=0.3
    )


def test_best_fit_leaves_a_deliberate_width_on_an_empty_column_alone(sheet, roundtrip):
    sheet["A1"] = "data"

    toolkit = WorksheetToolkit(sheet)
    toolkit.set_column_width(columns=3, width=30)
    toolkit.set_column_best_fit(columns=[1, 3])

    ws = roundtrip(sheet)
    assert ws.column_dimensions["C"].width == 30


def test_best_fit_sizes_a_date_from_its_displayed_form(sheet, roundtrip):
    sheet["A1"] = datetime(2026, 9, 10)
    sheet["A1"].number_format = "yyyy-mm-dd"
    sheet["B1"] = "2026-09-10"

    WorksheetToolkit(sheet).set_column_best_fit(columns=[1, 2])

    ws = roundtrip(sheet)
    assert ws.column_dimensions["A"].width == pytest.approx(
        ws.column_dimensions["B"].width, abs=1.0
    )


def test_column_width_is_clamped_to_the_excel_maximum(sheet, roundtrip):
    WorksheetToolkit(sheet).set_column_width(columns=1, width=300)

    assert roundtrip(sheet).column_dimensions["A"].width == EXCEL_MAX_COLUMN_WIDTH


def test_zero_column_width_collapses_the_column_in_the_saved_file(sheet, roundtrip):
    WorksheetToolkit(sheet).set_column_width(columns=1, width=0)

    dimension = roundtrip(sheet).column_dimensions["A"]
    assert dimension.width == 0 or dimension.hidden


def test_zero_row_height_collapses_the_row_in_the_saved_file(sheet, roundtrip):
    WorksheetToolkit(sheet).set_row_height(rows=1, height=0)

    dimension = roundtrip(sheet).row_dimensions[1]
    assert dimension.height == 0 or dimension.hidden


def test_negative_column_width_is_rejected(sheet):
    with pytest.raises((ValueError, TypeError)):
        WorksheetToolkit(sheet).set_column_width(columns=1, width=-5)


def test_negative_row_height_is_rejected(sheet):
    with pytest.raises((ValueError, TypeError)):
        WorksheetToolkit(sheet).set_row_height(rows=1, height=-5)


def test_row_zero_is_rejected(sheet):
    with pytest.raises((ValueError, TypeError)):
        WorksheetToolkit(sheet).set_row_height(rows=0, height=20)


def test_column_width_without_a_width_is_rejected(sheet):
    with pytest.raises((TypeError, ValueError)):
        WorksheetToolkit(sheet).set_column_width(columns=1)


def test_row_height_without_a_height_is_rejected(sheet):
    with pytest.raises((TypeError, ValueError)):
        WorksheetToolkit(sheet).set_row_height(rows=1)


def test_best_fit_measures_an_unstyled_cell_at_the_workbook_default_size(sheet, roundtrip):
    """A cell with no font of its own inherits the workbook's size, not a hardcoded 11."""
    sheet.parent._fonts[0] = Font(name="Calibri", sz=22)
    sheet["A1"] = "hello world"
    sheet["B1"] = "hello world"
    sheet["B1"].font = Font(sz=22)
    sheet["C1"] = "hello world"
    sheet["C1"].font = Font(sz=11)

    WorksheetToolkit(sheet).set_column_best_fit(columns=[1, 2, 3])

    ws = roundtrip(sheet)
    inherited = ws.column_dimensions["A"].width
    assert inherited == pytest.approx(ws.column_dimensions["B"].width)
    assert inherited > ws.column_dimensions["C"].width


def test_best_fit_widens_for_a_cell_font_larger_than_the_workbook_default(sheet, roundtrip):
    """The width unit is the workbook's font, so a bigger cell font needs more units."""
    sheet.parent._fonts[0] = Font(name="Calibri", sz=11)
    sheet["A1"] = "hello world"
    sheet["B1"] = "hello world"
    sheet["B1"].font = Font(sz=22)

    WorksheetToolkit(sheet).set_column_best_fit(columns=[1, 2])

    ws = roundtrip(sheet)
    assert ws.column_dimensions["B"].width > ws.column_dimensions["A"].width * 1.8


def test_best_fit_renders_a_time_in_twelve_hour_form_when_the_format_says_so(sheet, roundtrip):
    """'2:07 PM' is shorter than '14:07:03', and shorter still than the stored value."""
    sheet["A1"] = datetime(2026, 9, 5, 14, 7, 3)
    sheet["A1"].number_format = "h:mm am/pm"
    sheet["B1"] = "2:07 PM"

    WorksheetToolkit(sheet).set_column_best_fit(columns=[1, 2])

    ws = roundtrip(sheet)
    assert ws.column_dimensions["A"].width == pytest.approx(
        ws.column_dimensions["B"].width, abs=1.0
    )


def test_best_fit_measures_a_general_formatted_value_as_stored(sheet, roundtrip):
    """Only dates are translated; anything else is measured as it is stored."""
    sheet["A1"] = 1234.5678
    sheet["B1"] = "1234.5678"

    WorksheetToolkit(sheet).set_column_best_fit(columns=[1, 2])

    ws = roundtrip(sheet)
    assert ws.column_dimensions["A"].width == pytest.approx(ws.column_dimensions["B"].width)


def test_row_height_is_clamped_to_the_excel_maximum(sheet, roundtrip):
    WorksheetToolkit(sheet).set_row_height(rows=1, height=500)

    assert roundtrip(sheet).row_dimensions[1].height == 409


def test_best_fit_matches_the_definition_of_a_column_width_unit(sheet, roundtrip):
    """A width of N means N digits fit, plus Excel's 5 pixels of cell padding."""
    for count in (1, 10, 20):
        sheet.cell(row=1, column=count, value="0" * count)

    WorksheetToolkit(sheet).set_column_best_fit(columns=[1, 10, 20])

    ws = roundtrip(sheet)
    for count, letter in ((1, "A"), (10, "J"), (20, "T")):
        assert ws.column_dimensions[letter].width == pytest.approx(count + 5 / 7, abs=0.01)


def test_best_fit_distinguishes_narrow_from_wide_characters(sheet, roundtrip):
    sheet["A1"] = "i" * 10
    sheet["B1"] = "W" * 10

    WorksheetToolkit(sheet).set_column_best_fit(columns=[1, 2])

    ws = roundtrip(sheet)
    assert ws.column_dimensions["B"].width > ws.column_dimensions["A"].width * 3


def test_best_fit_honours_min_and_max_width(sheet, roundtrip):
    sheet["A1"] = "x"
    sheet["B1"] = "a very long piece of text indeed"

    WorksheetToolkit(sheet).set_column_best_fit(columns=[1, 2], min_width=12, max_width=20)

    ws = roundtrip(sheet)
    assert ws.column_dimensions["A"].width == 12
    assert ws.column_dimensions["B"].width == 20


def test_best_fit_accepts_a_custom_measure(sheet, roundtrip):
    """A caller with metrics for another font supplies its own measurement."""
    sheet["A1"] = "hello"

    WorksheetToolkit(sheet).set_column_best_fit(
        columns=1, measure=lambda text, font, normal_font: len(text) * 2
    )

    assert roundtrip(sheet).column_dimensions["A"].width == 10


def test_best_fit_uses_the_metrics_of_the_cell_font(sheet, roundtrip):
    """Courier New is monospaced and wider than Calibri, so the same text needs more room."""
    sheet["A1"] = "iiiiiiiiii"
    sheet["B1"] = "iiiiiiiiii"
    sheet["B1"].font = Font(name="Courier New")

    WorksheetToolkit(sheet).set_column_best_fit(columns=[1, 2])

    ws = roundtrip(sheet)
    assert ws.column_dimensions["B"].width > ws.column_dimensions["A"].width * 1.5


@pytest.mark.parametrize(
    ("number_format", "shown"),
    [
        ("yyyy-mm-dd", "2026-09-05"),
        ("mm/dd/yyyy", "09/05/2026"),
        ("d/m/yyyy", "5/9/2026"),
        ("mmm d, yyyy", "Sep 5, 2026"),
        ("mmmm yyyy", "September 2026"),
        ("dddd", "Saturday"),
        ("hh:mm:ss", "14:07:03"),
        ("h:mm am/pm", "2:07 PM"),
        ("m/d/yy h:mm", "9/5/26 14:07"),
    ],
)
def test_a_date_is_measured_as_its_number_format_renders_it(sheet, roundtrip, number_format, shown):
    """The width must match the displayed text, whatever platform this runs on.

    The month-name tokens are the ones that break if a bare mm is matched first,
    and the unpadded hour is the one that breaks on Windows, where strftime has no
    %-I. Both are covered here rather than left to a manual check.
    """
    sheet["A1"] = datetime(2026, 9, 5, 14, 7, 3)
    sheet["A1"].number_format = number_format
    sheet["B1"] = shown

    WorksheetToolkit(sheet).set_column_best_fit(columns=[1, 2])

    ws = roundtrip(sheet)
    assert ws.column_dimensions["A"].width == pytest.approx(ws.column_dimensions["B"].width)
