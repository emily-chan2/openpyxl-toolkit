"""Behavior of ``set_border`` and ``set_outside_border``.

Everything is asserted on the reloaded worksheet: a border that only exists in
memory is not a border.
"""

import pytest
from openpyxl.styles import Border, Side

from openpyxl_toolkit import WorksheetToolkit


def test_thin_border_survives_on_the_named_sides(sheet, roundtrip):
    WorksheetToolkit(sheet).set_border(rows=[1], columns=[1], sides=("top", "bottom"), style="thin")

    border = roundtrip(sheet).cell(row=1, column=1).border

    assert (border.top.style, border.bottom.style) == ("thin", "thin")


def test_sides_that_were_not_named_stay_without_a_border(sheet, roundtrip):
    WorksheetToolkit(sheet).set_border(rows=[1], columns=[1], sides=("top", "bottom"), style="thin")

    border = roundtrip(sheet).cell(row=1, column=1).border

    assert (border.left.style, border.right.style) == (None, None)


def test_a_later_call_on_another_side_preserves_the_earlier_side(sheet, roundtrip):
    """The merge guarantee: a second call adds a side, it does not start over."""
    toolkit = WorksheetToolkit(sheet)
    toolkit.set_border(rows=[1], columns=[1], sides=("top",), style="thin")
    toolkit.set_border(rows=[1], columns=[1], sides=("left",), style="thick")

    border = roundtrip(sheet).cell(row=1, column=1).border

    assert (border.top.style, border.left.style) == ("thin", "thick")


def test_recoloring_a_side_leaves_its_style_intact(sheet, roundtrip):
    """Passing only ``color`` must not drop the style set by an earlier call."""
    toolkit = WorksheetToolkit(sheet)
    toolkit.set_border(rows=[1], columns=[1], sides=("top",), style="medium")
    toolkit.set_border(rows=[1], columns=[1], sides=("top",), color="#0000ff")

    border = roundtrip(sheet).cell(row=1, column=1).border

    assert border.top.style == "medium"


def test_applying_a_border_leaves_an_existing_font_alone(sheet, roundtrip):
    toolkit = WorksheetToolkit(sheet)
    toolkit.set_font(rows=[1], columns=[1], bold=True)
    toolkit.set_border(rows=[1], columns=[1], sides=("top",), style="thin")

    assert roundtrip(sheet).cell(row=1, column=1).font.bold is True


def test_outside_border_draws_the_perimeter_of_the_block(sheet, roundtrip):
    WorksheetToolkit(sheet).set_outside_border(
        start_row=1, end_row=3, start_column=1, end_column=3, style="thin"
    )

    reloaded = roundtrip(sheet)
    perimeter = (
        reloaded.cell(row=1, column=2).border.top.style,
        reloaded.cell(row=3, column=2).border.bottom.style,
        reloaded.cell(row=2, column=1).border.left.style,
        reloaded.cell(row=2, column=3).border.right.style,
    )

    assert perimeter == ("thin", "thin", "thin", "thin")


def test_outside_border_leaves_an_interior_cell_unbordered(sheet, roundtrip):
    WorksheetToolkit(sheet).set_outside_border(
        start_row=1, end_row=3, start_column=1, end_column=3, style="thin"
    )

    border = roundtrip(sheet).cell(row=2, column=2).border
    interior = (
        border.top.style,
        border.bottom.style,
        border.left.style,
        border.right.style,
    )

    assert interior == (None, None, None, None)


def test_outside_border_leaves_the_inward_sides_of_an_edge_cell_clear(sheet, roundtrip):
    """A cell on the top edge gets a top border only, not a box around it."""
    WorksheetToolkit(sheet).set_outside_border(
        start_row=1, end_row=3, start_column=1, end_column=3, style="thin"
    )

    border = roundtrip(sheet).cell(row=1, column=2).border
    inward = (border.bottom.style, border.left.style, border.right.style)

    assert inward == (None, None, None)


def test_border_color_persists_as_the_same_argb_as_font_color(sheet, roundtrip):
    """The same hex handed to two methods must round trip to one ARGB string."""
    toolkit = WorksheetToolkit(sheet)
    toolkit.set_border(rows=[1], columns=[1], sides=("top",), style="thin", color="#ff0000")
    toolkit.set_font(rows=[1], columns=[2], color="#ff0000")

    reloaded = roundtrip(sheet)

    assert (
        reloaded.cell(row=1, column=1).border.top.color.rgb
        == reloaded.cell(row=1, column=2).font.color.rgb
    )


def test_a_color_with_no_style_is_rejected(sheet):
    """A side with no style draws nothing, so a color alone cannot show."""
    with pytest.raises(ValueError, match="style='thin'"):
        WorksheetToolkit(sheet).set_border(rows=[1], columns=[1], sides=("top",), color="#ff0000")


def test_a_color_with_no_style_is_fine_when_the_side_already_exists(sheet, roundtrip):
    sheet.cell(row=1, column=1, value="x").border = Border(top=Side(style="thin"))

    WorksheetToolkit(sheet).set_border(rows=[1], columns=[1], sides=("top",), color="#ff0000")

    top = roundtrip(sheet).cell(row=1, column=1).border.top
    assert top.style == "thin" and top.color.rgb.lower().endswith("ff0000")


def test_a_style_with_no_color_is_allowed(sheet, roundtrip):
    """An uncolored border is drawn in Excel's automatic color, which is wanted."""
    WorksheetToolkit(sheet).set_border(rows=[1], columns=[1], sides=("top",), style="thin")

    assert roundtrip(sheet).cell(row=1, column=1).border.top.style == "thin"


def test_a_rejected_border_leaves_the_whole_range_untouched(sheet, roundtrip):
    for c in (1, 2, 3):
        sheet.cell(row=1, column=c, value="x")
    sheet["A1"].border = Border(top=Side(style="thin"))

    with pytest.raises(ValueError):
        WorksheetToolkit(sheet).set_border(rows=[1], sides=("top",), color="#ff0000")

    assert roundtrip(sheet)["A1"].border.top.color is None


def test_diagonal_up_and_diagonal_down_can_coexist(sheet, roundtrip):
    toolkit = WorksheetToolkit(sheet)
    toolkit.set_border(rows=[1], columns=[1], sides=("diagonal_up",), style="thin")
    toolkit.set_border(rows=[1], columns=[1], sides=("diagonal_down",), style="thin")

    border = roundtrip(sheet).cell(row=1, column=1).border

    assert (border.diagonalUp, border.diagonalDown) == (True, True)


def test_a_diagonal_is_taken_away_by_a_null_style(sheet, roundtrip):
    toolkit = WorksheetToolkit(sheet)
    toolkit.set_border(rows=[1], columns=[1], sides=("diagonal_up",), style="thin")
    toolkit.set_border(rows=[1], columns=[1], sides=("diagonal_up",), style=None)

    border = roundtrip(sheet).cell(row=1, column=1).border

    assert (border.diagonalUp, border.diagonalDown) == (False, False)
    assert border.diagonal.style is None


def test_taking_one_diagonal_away_leaves_the_other_drawn(sheet, roundtrip):
    """The two share one line, so removing one must not rub out the other."""
    toolkit = WorksheetToolkit(sheet)
    toolkit.set_border(rows=[1], columns=[1], sides=("diagonal_up", "diagonal_down"), style="thin")
    toolkit.set_border(rows=[1], columns=[1], sides=("diagonal_up",), style=None)

    border = roundtrip(sheet).cell(row=1, column=1).border

    assert (border.diagonalUp, border.diagonalDown) == (False, True)
    assert border.diagonal.style == "thin"


def test_the_shared_line_goes_once_neither_diagonal_is_drawn(sheet, roundtrip):
    toolkit = WorksheetToolkit(sheet)
    toolkit.set_border(rows=[1], columns=[1], sides=("diagonal_up", "diagonal_down"), style="thin")
    toolkit.set_border(rows=[1], columns=[1], sides=("diagonal_up", "diagonal_down"), style=None)

    border = roundtrip(sheet).cell(row=1, column=1).border

    assert (border.diagonalUp, border.diagonalDown) == (False, False)
    assert border.diagonal.style is None


def test_naming_a_diagonal_with_no_style_draws_nothing(sheet, roundtrip):
    """A flag with no line behind it claims a diagonal the cell does not show."""
    WorksheetToolkit(sheet).set_border(rows=[1], columns=[1], sides=("diagonal_up",))

    border = roundtrip(sheet).cell(row=1, column=1).border

    assert border.diagonalUp is False


def test_taking_a_diagonal_away_leaves_the_straight_sides_alone(sheet, roundtrip):
    toolkit = WorksheetToolkit(sheet)
    toolkit.set_border(rows=[1], columns=[1], sides=("top", "left"), style="medium")
    toolkit.set_border(rows=[1], columns=[1], sides=("diagonal_down",), style="thin")
    toolkit.set_border(rows=[1], columns=[1], sides=("diagonal_down",), style=None)

    border = roundtrip(sheet).cell(row=1, column=1).border

    assert (border.top.style, border.left.style) == ("medium", "medium")
    assert border.diagonalDown is False


def test_a_color_cannot_be_given_while_taking_a_diagonal_away(sheet):
    toolkit = WorksheetToolkit(sheet)
    toolkit.set_border(rows=[1], columns=[1], sides=("diagonal_up",), style="thin")

    with pytest.raises(ValueError):
        toolkit.set_border(
            rows=[1], columns=[1], sides=("diagonal_up",), style=None, color="#ff0000"
        )


def test_an_unknown_side_name_is_rejected(sheet):
    with pytest.raises(ValueError):
        WorksheetToolkit(sheet).set_border(rows=[1], columns=[1], sides=("lft",), style="thin")


def test_outside_border_without_a_style_is_rejected(sheet):
    """style has no default: a border with no style draws nothing."""
    with pytest.raises(TypeError):
        WorksheetToolkit(sheet).set_outside_border(
            start_row=1, end_row=3, start_column=1, end_column=3
        )
