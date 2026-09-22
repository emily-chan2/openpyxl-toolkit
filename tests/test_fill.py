"""Behavior of ``WorksheetToolkit.set_fill``, asserted on the reloaded workbook."""

import pytest
from openpyxl.styles import Color, GradientFill, PatternFill

from openpyxl_toolkit import WorksheetToolkit


def _grid(ws, rows=3, columns=3):
    for row in range(1, rows + 1):
        for column in range(1, columns + 1):
            ws.cell(row=row, column=column, value=f"r{row}c{column}")
    return ws


def _painted(ws):
    """Coordinates of every cell carrying a pattern of any kind."""
    return {
        ws.cell(row=row, column=column).coordinate
        for row in range(1, ws.max_row + 1)
        for column in range(1, ws.max_column + 1)
        if ws.cell(row=row, column=column).fill.fill_type is not None
    }


def test_solid_fill_survives_the_round_trip(sheet, roundtrip):
    _grid(sheet)
    WorksheetToolkit(sheet).set_fill(
        rows=[1], columns=[1], fill_type="solid", start_color="#f4d2d3"
    )

    fill = roundtrip(sheet)["A1"].fill
    assert (fill.fill_type, fill.start_color.rgb.upper()) == ("solid", "FFF4D2D3")


def test_fill_type_survives_a_later_color_only_call(sheet, roundtrip):
    """Changing the color must not throw away the pattern set by an earlier call."""
    _grid(sheet)
    toolkit = WorksheetToolkit(sheet)
    toolkit.set_fill(rows=[1], columns=[1], fill_type="solid", start_color="#ff0000")
    toolkit.set_fill(rows=[1], columns=[1], start_color="#00ff00")

    fill = roundtrip(sheet)["A1"].fill
    assert (fill.fill_type, fill.start_color.rgb.upper()) == ("solid", "FF00FF00")


def test_start_color_survives_a_later_fill_type_only_call(sheet, roundtrip):
    """Changing the pattern must not throw away the color set by an earlier call."""
    _grid(sheet)
    toolkit = WorksheetToolkit(sheet)
    toolkit.set_fill(rows=[1], columns=[1], fill_type="solid", start_color="#ff0000")
    toolkit.set_fill(rows=[1], columns=[1], fill_type="lightGrid")

    fill = roundtrip(sheet)["A1"].fill
    assert (fill.fill_type, fill.start_color.rgb.upper()) == ("lightGrid", "FFFF0000")


def test_end_color_survives_a_later_start_color_call(sheet, roundtrip):
    """The pattern background is the third attribute that must be merged, not reset."""
    _grid(sheet)
    toolkit = WorksheetToolkit(sheet)
    toolkit.set_fill(
        rows=[1], columns=[1], fill_type="lightGrid", start_color="#ff0000", end_color="#0000ff"
    )
    toolkit.set_fill(rows=[1], columns=[1], start_color="#00ff00")

    fill = roundtrip(sheet)["A1"].fill
    assert fill.end_color.rgb.upper() == "FF0000FF"


def test_intersections_only_paints_the_intersection(sheet, roundtrip):
    _grid(sheet)
    WorksheetToolkit(sheet).set_fill(
        rows=[1, 2],
        columns=[1, 3],
        intersections_only=True,
        fill_type="solid",
        start_color="#ff0000",
    )

    assert _painted(roundtrip(sheet)) == {"A1", "A2", "C1", "C2"}


def test_union_selection_paints_whole_rows_and_whole_columns(sheet, roundtrip):
    _grid(sheet)
    WorksheetToolkit(sheet).set_fill(
        rows=[1], columns=[1], intersections_only=False, fill_type="solid", start_color="#ff0000"
    )

    assert _painted(roundtrip(sheet)) == {"A1", "B1", "C1", "A2", "A3"}


def test_row_selection_without_columns_spans_the_used_range(sheet, roundtrip):
    _grid(sheet)
    WorksheetToolkit(sheet).set_fill(rows=[2], fill_type="solid", start_color="#ff0000")

    assert _painted(roundtrip(sheet)) == {"A2", "B2", "C2"}


def test_empty_selection_paints_every_used_cell(sheet, roundtrip):
    _grid(sheet, rows=2, columns=2)
    WorksheetToolkit(sheet).set_fill(fill_type="solid", start_color="#ff0000")

    assert _painted(roundtrip(sheet)) == {"A1", "A2", "B1", "B2"}


def test_existing_theme_color_is_kept_when_only_the_pattern_changes(sheet, roundtrip):
    _grid(sheet)
    sheet["A1"].fill = PatternFill(
        patternType="solid", fgColor=Color(theme=9, tint=0.6), bgColor=Color(indexed=64)
    )

    WorksheetToolkit(sheet).set_fill(rows=[1], columns=[1], fill_type="solid")

    color = roundtrip(sheet)["A1"].fill.start_color
    assert (color.type, color.theme, color.tint) == ("theme", 9, pytest.approx(0.6))


def test_existing_indexed_color_is_kept_when_only_the_pattern_changes(sheet, roundtrip):
    _grid(sheet)
    sheet["A1"].fill = PatternFill(
        patternType="solid", fgColor=Color(indexed=5), bgColor=Color(indexed=64)
    )

    WorksheetToolkit(sheet).set_fill(rows=[1], columns=[1], fill_type="solid")

    color = roundtrip(sheet)["A1"].fill.start_color
    assert (color.type, color.indexed) == ("indexed", 5)


def test_existing_automatic_color_is_kept_when_only_the_pattern_changes(sheet, roundtrip):
    _grid(sheet)
    sheet["A1"].fill = PatternFill(
        patternType="solid", fgColor=Color(auto=True), bgColor=Color(indexed=64)
    )

    WorksheetToolkit(sheet).set_fill(rows=[1], columns=[1], fill_type="solid")

    color = roundtrip(sheet)["A1"].fill.start_color
    assert (color.type, color.auto) == ("auto", True)


def test_a_gradient_filled_cell_can_be_given_a_solid_fill(sheet, roundtrip):
    _grid(sheet)
    sheet["A1"].fill = GradientFill(stop=("FFFFFFFF", "FFFF0000"))

    WorksheetToolkit(sheet).set_fill(
        rows=[1], columns=[1], fill_type="solid", start_color="#00ff00"
    )

    fill = roundtrip(sheet)["A1"].fill
    assert (fill.fill_type, fill.start_color.rgb.upper()) == ("solid", "FF00FF00")


def test_a_color_on_its_own_is_rejected(sheet):
    """A cell with no pattern cannot show a color, so the call says so instead."""
    _grid(sheet)
    with pytest.raises(ValueError, match="fill_type='solid'"):
        WorksheetToolkit(sheet).set_fill(rows=[1], start_color="#f4d2d3")


def test_a_color_on_its_own_is_fine_when_the_cell_is_already_filled(sheet, roundtrip):
    _grid(sheet)
    sheet["A1"].fill = PatternFill("solid", start_color="FF00FF00")

    WorksheetToolkit(sheet).set_fill(rows=[1], columns=[1], start_color="#f4d2d3")

    assert roundtrip(sheet)["A1"].fill.start_color.rgb.upper() == "FFF4D2D3"


def test_a_pattern_with_no_color_is_rejected(sheet):
    """A pattern with no color paints the cell black, so it is rejected."""
    _grid(sheet)
    with pytest.raises(ValueError, match="start_color"):
        WorksheetToolkit(sheet).set_fill(rows=[1], columns=[1], fill_type="solid")


def test_a_pattern_with_no_color_is_fine_when_the_cell_is_already_colored(sheet, roundtrip):
    _grid(sheet)
    sheet["A1"].fill = PatternFill("solid", start_color="FF00FF00")

    WorksheetToolkit(sheet).set_fill(rows=[1], columns=[1], fill_type="lightGrid")

    reloaded = roundtrip(sheet)["A1"].fill
    assert (reloaded.fill_type, reloaded.start_color.rgb) == ("lightGrid", "FF00FF00")


def test_removing_a_fill_is_still_allowed(sheet, roundtrip):
    """fill_type=None means remove it, and must not trip the color checks."""
    _grid(sheet)
    sheet["A1"].fill = PatternFill("solid", start_color="FF00FF00")

    WorksheetToolkit(sheet).set_fill(rows=[1], columns=[1], fill_type=None)

    assert roundtrip(sheet)["A1"].fill.fill_type is None


def test_a_failing_call_leaves_no_cell_in_the_range_painted(sheet, roundtrip, monkeypatch):
    """A range is formatted as a unit: if one cell cannot be filled, none is."""
    _grid(sheet)
    attempts = []

    def fail_on_the_third_cell(*args, **kwargs):
        attempts.append(None)
        if len(attempts) == 3:
            raise RuntimeError("cell could not be filled")
        return PatternFill(*args, **kwargs)

    monkeypatch.setattr("openpyxl_toolkit.worksheet_toolkit.PatternFill", fail_on_the_third_cell)

    with pytest.raises(RuntimeError):
        WorksheetToolkit(sheet).set_fill(
            rows=[1], columns=[1, 2, 3], fill_type="solid", start_color="#ff0000"
        )

    assert _painted(roundtrip(sheet)) == set()


def test_a_cell_does_not_share_a_color_object_with_the_fill_it_was_read_from(sheet, roundtrip):
    """A Color is stored by reference; sharing one lets a mutation repaint cells."""
    _grid(sheet)
    header = PatternFill(patternType="solid", fgColor=Color(rgb="FF3366CC"))
    sheet["A1"].fill = header

    WorksheetToolkit(sheet).set_fill(rows=[1], columns=[1], fill_type="lightGrid")
    header.start_color.rgb = "FFFF0000"

    assert roundtrip(sheet)["A1"].fill.start_color.rgb == "FF3366CC"


def test_a_gradient_survives_a_call_that_asks_for_nothing(sheet, roundtrip):
    """No style arguments must leave a gradient alone, not flatten it."""
    _grid(sheet)
    sheet["A1"].fill = GradientFill(stop=("FFFFFFFF", "FFFF0000"))

    WorksheetToolkit(sheet).set_fill(rows=[1], columns=[1])

    assert roundtrip(sheet)["A1"].fill.tagname == "gradientFill"


def test_clearing_a_color_is_rejected_rather_than_silently_repainting(sheet):
    """A pattern fill has no None state: its default foreground is opaque black."""
    _grid(sheet)
    toolkit = WorksheetToolkit(sheet)
    for kwargs in ({"start_color": None}, {"end_color": None}):
        with pytest.raises(ValueError, match="cannot clear a color"):
            toolkit.set_fill(rows=[1], columns=[1], **kwargs)
